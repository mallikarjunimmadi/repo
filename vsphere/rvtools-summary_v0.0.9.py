#!/usr/bin/env python3

from __future__ import annotations

import argparse
import logging
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

try:
    import pandas as pd
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "This script requires pandas. "
        "Run it with the Python environment where that package is installed, "
        "for example: ../vsphere/bin/python3 rvtools-summary.py [input_dir]"
    ) from exc

try:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover
    Font = None
    get_column_letter = None
    OPENPYXL_AVAILABLE = False


SCRIPT_VERSION = "0.0.9"
MANDATORY_TAG_COLUMNS = [
    "Application",
    "APMID",
    "Department",
    "Tier",
    "Location",
    "Environment",
    "Criticality",
    "AppMode",
]

LATEST_HW_VERSION = "21"
GENERATED_REPORT_BASENAMES = {
    "rvtools-summary-report.xlsx",
    "rvtools-summary-report.csv",
}
COMBINED_INVENTORY_SHEETS = [
    "vHost",
    "vHBA",
    "vNIC",
]
CONSOLE_LOG_BASENAME = "rvtools-summary.log"
EXCEL_SHEET_NAME_ALIASES = {
    "vms_not_using_paravirtual_scsi_adapter": "vms_not_using_pvscsi",
}
SUMMARY_HEADER_SHEET_LINKS = {
    "ds_cross_clustered": "ds_cross_clustered",
    "ds_with_no_vms": "ds_with_no_vms",
    "vms_powered_off": "vms_powered_off",
    "vms_suspended": "vms_suspended",
    "vms_templates": "templates",
    "vms_consolidation_needed": "consolidation_needed",
    "vms_in_not_connected_state": "vms_not_in_connected_state",
    "vms_cpu_socket_fix_needed": "vms_cpu_socket_fix_needed",
    "vms_with_cpu_reservation": "vms_with_cpu_reservation",
    "vms_with_memory_reservation": "vms_with_memory_reservation",
    "vms_with_memory_contention": "vms_with_memory_contention",
    "vms_not_using_paravirtual_scsi_adapter": "vms_not_using_paravirtual_scsi_adapter",
    "vms_network_not_vmxnet3": "vms_network_not_vmxnet3",
    "vms_with_duplicate_ip": "vms_with_duplicate_ip",
    "vms_with_duplicate_mac_address": "vms_with_duplicate_mac_address",
    "duplicate_ips": "duplicate_ips",
    "duplicate_mac": "duplicate_mac",
    "host_with_duplicate_ip": "host_with_duplicate_ip",
    "host_with_duplicate_mac": "host_with_duplicate_mac",
    "vms_with_usb": "vms_with_usb",
    "vms_with_snapshot": "vms_with_snapshots",
    "vms_hot_add_enabled": "vms_hot_add_enabled",
    "vms_tools_fix_needed": "vms_tools_fix_needed",
    "clusters_total": "cluster_configuration_issues",
    "clusters_ha_disabled": "cluster_configuration_issues",
    "clusters_drs_disabled": "cluster_configuration_issues",
    "clusters_drs_not_fullyautomatic": "cluster_configuration_issues",
    "host_isolation_response_not_none": "cluster_configuration_issues",
    "host_cpu_models": "host_cpu_models",
    "host_cpu_overprovisioned": "host_cpu_overprovisioned",
    "host_in_mm": "host_in_mm",
    "host_ht_not_active": "host_ht_not_active",
    "host_mem_pressure": "host_mem_pressure",
    "host_power_policy_not_high_perf": "host_power_policy_not_high_perf",
    "host_esxi_versions": "host_esxi_versions",
    "host_ntp_not_running": "host_ntp_not_running",
    "host_vendors": "host_vendors",
    "host_models": "host_models",
    "datastores_with_extents_gt_1": "datastores_with_extents_gt_1",
    "vms_hw_version_not_latest": "hw_version_not_latest",
    "vms_missing_mandatory_tags": "missing_mandatory_tags",
    "vms_with_duplicate_vm_uuid": "vms_with_duplicate_vm_uuid",
}
MIB_PER_GIB = 1024
MIB_PER_TB = 1024 * 1024
HOST_CPU_OVERPROVISIONED_VCPUS_PER_CORE_THRESHOLD = 2.0

# Avi/NSX ALB CSV report exports carry a timestamp suffix
# (e.g. "serviceengine_report_20260709_083522.csv"), so reports are matched
# by substring rather than exact filename. Order matters: "vsvip_config_report"
# is checked before "vs_config_report" so it can't be shadowed by it (it
# isn't a substring match today, but keep the more specific pattern first).
AVI_REPORT_FILE_PATTERNS: dict[str, str] = {
    "service_engines": "serviceengine_report",
    "vsvips": "vsvip_config_report",
    "virtual_services": "vs_config_report",
    "network_subnets": "network_subnet_report",
}


@dataclass
class WorkbookSummary:
    path: Path
    source_vcenter: str = "unknown"
    vm_total: int = 0
    datastore_total: int = 0
    vsan_datastore_total: int = 0
    vsan_datastore_with_no_vms: int = 0
    vsan_datastore_capacity_mib: float = 0.0
    ds_cross_clustered: int = 0
    ds_with_no_vms: int = 0
    cluster_total: int = 0
    clusters_ha_disabled: int = 0
    clusters_drs_disabled: int = 0
    clusters_drs_not_fullyautomatic: int = 0
    host_isolation_response_not_none: int = 0
    host_total: int = 0
    host_in_mm: int = 0
    host_ht_not_active: int = 0
    host_cpu_overprovisioned: int = 0
    host_mem_pressure: int = 0
    host_power_policy_not_high_perf: int = 0
    host_esxi_versions: int = 0
    host_ntp_not_running: int = 0
    host_vendors: int = 0
    host_models: int = 0
    host_cpu_cores: int = 0
    host_memory_mib: float = 0.0
    powered_on_vms: int = 0
    powered_off_vms: int = 0
    suspended_vms: int = 0
    template_vms: int = 0
    srm_placeholder_vms: int = 0
    consolidation_needed_vms: int = 0
    vms_not_connected: int = 0
    vms_cpu_socket_fix_needed: int = 0
    vms_with_cpu_reservation: int = 0
    vms_with_memory_reservation: int = 0
    vms_mem_ballooned: int = 0
    vms_mem_swapped: int = 0
    vms_with_memory_contention: int = 0
    vms_not_using_paravirtual_scsi_adapter: int = 0
    vms_network_not_vmxnet3: int = 0
    vms_with_duplicate_ip: int = 0
    vms_with_duplicate_mac_address: int = 0
    host_with_duplicate_ip: int = 0
    host_with_duplicate_mac: int = 0
    duplicate_ips: int = 0
    duplicate_mac: int = 0
    vms_with_usb: int = 0
    vms_with_snapshot: int = 0
    vms_with_snapshot_age_le_48h: int = 0
    vms_with_snapshot_age_gt_48h_le_7d: int = 0
    vms_with_snapshot_age_gt_7d: int = 0
    vms_hot_add_enabled: int = 0
    vms_tools_fix_needed: int = 0
    datastores_with_extents_gt_1: int = 0
    hw_version_not_latest_vms: int = 0
    duplicate_uuid_vms: int = 0
    vms_missing_mandatory_tags: int = 0
    cpu_le_8: int = 0
    cpu_gt_8_le_16: int = 0
    cpu_gt_16_le_32: int = 0
    cpu_gt_32_le_64: int = 0
    cpu_gt_64: int = 0
    memory_le_16_gib: int = 0
    memory_gt_16_le_32_gib: int = 0
    memory_gt_32_le_64_gib: int = 0
    memory_gt_64_le_128_gib: int = 0
    memory_gt_128_gib: int = 0
    vm_vcpu_total: int = 0
    vm_memory_mib: float = 0.0
    provisioned_mib: float = 0.0
    used_mib: float = 0.0
    poweroff_provisioned_mib: float = 0.0
    poweroff_used_mib: float = 0.0
    power_states: Counter[str] = field(default_factory=Counter)
    connection_states: Counter[str] = field(default_factory=Counter)
    efi_secure_boot: Counter[str] = field(default_factory=Counter)
    hw_versions: Counter[str] = field(default_factory=Counter)
    missing_tag_counts: Counter[str] = field(default_factory=Counter)
    vm_uuid_counts: Counter[str] = field(default_factory=Counter)
    host_cpu_model_counts: Counter[str] = field(default_factory=Counter)
    host_esxi_version_counts: Counter[str] = field(default_factory=Counter)
    host_vendor_counts: Counter[str] = field(default_factory=Counter)
    host_model_counts: Counter[str] = field(default_factory=Counter)
    controller_type_counts: Counter[str] = field(default_factory=Counter)
    network_adapter_type_counts: Counter[str] = field(default_factory=Counter)
    # Raw per-workbook contributions used to resolve duplicate IP/MAC
    # addresses globally, across all workbooks together (see
    # apply_global_duplicate_network_counts). Each Counter maps a value
    # (ip/mac) to how many distinct entities of that type, in this workbook,
    # reported it.
    host_ip_counts: Counter[str] = field(default_factory=Counter)
    host_mac_counts: Counter[str] = field(default_factory=Counter)
    vm_ip_counts: Counter[str] = field(default_factory=Counter)
    vm_mac_counts: Counter[str] = field(default_factory=Counter)

    def merge(self, other: "WorkbookSummary") -> None:
        self.vm_total += other.vm_total
        self.datastore_total += other.datastore_total
        self.vsan_datastore_total += other.vsan_datastore_total
        self.vsan_datastore_with_no_vms += other.vsan_datastore_with_no_vms
        self.vsan_datastore_capacity_mib += other.vsan_datastore_capacity_mib
        self.ds_cross_clustered += other.ds_cross_clustered
        self.ds_with_no_vms += other.ds_with_no_vms
        self.cluster_total += other.cluster_total
        self.clusters_ha_disabled += other.clusters_ha_disabled
        self.clusters_drs_disabled += other.clusters_drs_disabled
        self.clusters_drs_not_fullyautomatic += other.clusters_drs_not_fullyautomatic
        self.host_isolation_response_not_none += other.host_isolation_response_not_none
        self.host_total += other.host_total
        self.host_in_mm += other.host_in_mm
        self.host_ht_not_active += other.host_ht_not_active
        self.host_cpu_overprovisioned += other.host_cpu_overprovisioned
        self.host_mem_pressure += other.host_mem_pressure
        self.host_power_policy_not_high_perf += other.host_power_policy_not_high_perf
        self.host_esxi_versions += other.host_esxi_versions
        self.host_ntp_not_running += other.host_ntp_not_running
        self.host_vendors += other.host_vendors
        self.host_models += other.host_models
        self.host_cpu_cores += other.host_cpu_cores
        self.host_memory_mib += other.host_memory_mib
        self.powered_on_vms += other.powered_on_vms
        self.powered_off_vms += other.powered_off_vms
        self.suspended_vms += other.suspended_vms
        self.template_vms += other.template_vms
        self.srm_placeholder_vms += other.srm_placeholder_vms
        self.consolidation_needed_vms += other.consolidation_needed_vms
        self.vms_not_connected += other.vms_not_connected
        self.vms_cpu_socket_fix_needed += other.vms_cpu_socket_fix_needed
        self.vms_with_cpu_reservation += other.vms_with_cpu_reservation
        self.vms_with_memory_reservation += other.vms_with_memory_reservation
        self.vms_mem_ballooned += other.vms_mem_ballooned
        self.vms_mem_swapped += other.vms_mem_swapped
        self.vms_with_memory_contention += other.vms_with_memory_contention
        self.vms_not_using_paravirtual_scsi_adapter += other.vms_not_using_paravirtual_scsi_adapter
        self.vms_network_not_vmxnet3 += other.vms_network_not_vmxnet3
        # vms_with_duplicate_ip, vms_with_duplicate_mac_address,
        # host_with_duplicate_ip, host_with_duplicate_mac, duplicate_ips and
        # duplicate_mac are intentionally not summed here -- they're
        # recomputed fresh from the merged host_ip_counts/vm_ip_counts/
        # host_mac_counts/vm_mac_counts below, since duplicate detection is
        # global (see build_combined_summary).
        self.vms_with_usb += other.vms_with_usb
        self.vms_with_snapshot += other.vms_with_snapshot
        self.vms_with_snapshot_age_le_48h += other.vms_with_snapshot_age_le_48h
        self.vms_with_snapshot_age_gt_48h_le_7d += other.vms_with_snapshot_age_gt_48h_le_7d
        self.vms_with_snapshot_age_gt_7d += other.vms_with_snapshot_age_gt_7d
        self.vms_hot_add_enabled += other.vms_hot_add_enabled
        self.vms_tools_fix_needed += other.vms_tools_fix_needed
        self.datastores_with_extents_gt_1 += other.datastores_with_extents_gt_1
        self.hw_version_not_latest_vms += other.hw_version_not_latest_vms
        self.duplicate_uuid_vms += other.duplicate_uuid_vms
        self.vms_missing_mandatory_tags += other.vms_missing_mandatory_tags
        self.cpu_le_8 += other.cpu_le_8
        self.cpu_gt_8_le_16 += other.cpu_gt_8_le_16
        self.cpu_gt_16_le_32 += other.cpu_gt_16_le_32
        self.cpu_gt_32_le_64 += other.cpu_gt_32_le_64
        self.cpu_gt_64 += other.cpu_gt_64
        self.memory_le_16_gib += other.memory_le_16_gib
        self.memory_gt_16_le_32_gib += other.memory_gt_16_le_32_gib
        self.memory_gt_32_le_64_gib += other.memory_gt_32_le_64_gib
        self.memory_gt_64_le_128_gib += other.memory_gt_64_le_128_gib
        self.memory_gt_128_gib += other.memory_gt_128_gib
        self.vm_vcpu_total += other.vm_vcpu_total
        self.vm_memory_mib += other.vm_memory_mib
        self.provisioned_mib += other.provisioned_mib
        self.used_mib += other.used_mib
        self.poweroff_provisioned_mib += other.poweroff_provisioned_mib
        self.poweroff_used_mib += other.poweroff_used_mib
        self.power_states.update(other.power_states)
        self.connection_states.update(other.connection_states)
        self.efi_secure_boot.update(other.efi_secure_boot)
        self.hw_versions.update(other.hw_versions)
        self.missing_tag_counts.update(other.missing_tag_counts)
        self.vm_uuid_counts.update(other.vm_uuid_counts)
        self.host_cpu_model_counts.update(other.host_cpu_model_counts)
        self.host_esxi_version_counts.update(other.host_esxi_version_counts)
        self.host_vendor_counts.update(other.host_vendor_counts)
        self.host_model_counts.update(other.host_model_counts)
        self.controller_type_counts.update(other.controller_type_counts)
        self.network_adapter_type_counts.update(other.network_adapter_type_counts)
        self.host_ip_counts.update(other.host_ip_counts)
        self.host_mac_counts.update(other.host_mac_counts)
        self.vm_ip_counts.update(other.vm_ip_counts)
        self.vm_mac_counts.update(other.vm_mac_counts)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Prepare a summary report from one or more RVTools XLSX exports. "
            "If no directory is supplied, the current working directory is used."
        )
    )
    parser.add_argument(
        "--version",
        action="version",
        version=f"%(prog)s {SCRIPT_VERSION}",
    )
    parser.add_argument(
        "-i",
        "--input-dir",
        default=".",
        help="Directory containing RVTools XLSX exports. Defaults to the current working directory.",
    )
    parser.add_argument(
        "-o",
        "--report-dir",
        default=".",
        help="Directory where the summary CSV and Excel reports will be written. Defaults to the current working directory.",
    )
    parser.add_argument(
        "--summary",
        action="store_true",
        help="Show per-vCenter summary on the console. By default only the combined/consolidated summary is shown.",
    )
    parser.add_argument(
        "--avi-data",
        dest="avi_data",
        default=None,
        help=(
            "Optional directory containing Avi/NSX Advanced Load Balancer CSV reports "
            "(service engine, VsVip, virtual service, network subnet exports). When supplied, "
            "Avi service engine and VIP IP/MAC addresses are folded into the duplicate IP/MAC "
            "detection alongside vSphere hosts and VMs."
        ),
    )
    return parser.parse_args()


def configure_logger(log_path: Path) -> logging.Logger:
    logger = logging.getLogger("rvtools_summary")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    logger.propagate = False

    class ConsoleOnlyFilter(logging.Filter):
        def filter(self, record: logging.LogRecord) -> bool:
            return getattr(record, "console", True)

    formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s")

    file_handler = logging.FileHandler(log_path, encoding="utf-8")
    file_handler.setLevel(logging.INFO)
    file_handler.setFormatter(formatter)

    console_handler = logging.StreamHandler(sys.stderr)
    console_handler.setLevel(logging.INFO)
    console_handler.setFormatter(logging.Formatter("%(message)s"))
    console_handler.addFilter(ConsoleOnlyFilter())

    logger.addHandler(file_handler)
    logger.addHandler(console_handler)
    return logger


def discover_workbooks(input_dir: Path) -> list[Path]:
    return sorted(
        path
        for path in input_dir.iterdir()
        if path.is_file()
        and path.suffix.lower() == ".xlsx"
        and not path.name.startswith("~$")
        and path.name not in GENERATED_REPORT_BASENAMES
    )


def load_sheet(workbook: Path, sheet_name: str) -> pd.DataFrame:
    return pd.read_excel(workbook, sheet_name=sheet_name)


def export_combined_inventory_csvs(workbooks: list[Path], output_dir: Path) -> list[Path]:
    if len(workbooks) <= 1:
        return []

    combined_csv_paths: list[Path] = []
    for sheet_name in COMBINED_INVENTORY_SHEETS:
        combined_frames: list[pd.DataFrame] = []
        for workbook_path in workbooks:
            workbook = pd.ExcelFile(workbook_path)
            if sheet_name not in workbook.sheet_names:
                continue
            frame = load_sheet(workbook_path, sheet_name).copy()
            frame.insert(0, "source_workbook", workbook_path.name)
            combined_frames.append(frame)

        if not combined_frames:
            continue

        combined_frame = pd.concat(combined_frames, ignore_index=True, sort=False)
        combined_csv_path = output_dir / f"rvtools-{sheet_name}-combined.csv"
        combined_frame.to_csv(combined_csv_path, index=False)
        combined_csv_paths.append(combined_csv_path)

    return combined_csv_paths


def discover_avi_report_files(avi_dir: Path) -> dict[str, list[Path]]:
    files_by_type: dict[str, list[Path]] = {report_type: [] for report_type in AVI_REPORT_FILE_PATTERNS}
    for path in sorted(avi_dir.iterdir()):
        if not (path.is_file() and path.suffix.lower() == ".csv"):
            continue
        for report_type, pattern in AVI_REPORT_FILE_PATTERNS.items():
            if pattern in path.name:
                files_by_type[report_type].append(path)
                break
    return files_by_type


def load_avi_reports(avi_dir: Path) -> dict[str, pd.DataFrame]:
    reports: dict[str, pd.DataFrame] = {}
    for report_type, paths in discover_avi_report_files(avi_dir).items():
        if not paths:
            continue
        frames = [pd.read_csv(path) for path in paths]
        reports[report_type] = pd.concat(frames, ignore_index=True, sort=False)
    return reports


def normalize_strings(series: pd.Series, default: str = "unknown") -> pd.Series:
    normalized = series.fillna("").astype(str).str.strip()
    return normalized.mask(normalized.eq(""), default)


def normalize_bool_strings(series: pd.Series) -> pd.Series:
    return normalize_strings(series, default="blank").str.lower()


def normalize_hw_version_series(series: pd.Series) -> pd.Series:
    """Normalize RVTools 'HW version' values to a stable string like "19".

    RVTools exports this column as numeric (e.g. 19). If even one row in the
    column is blank, pandas silently promotes the whole column to float64,
    and a naive str()/astype(str) turns every value into "19.0" instead of
    "19" -- which would then never match LATEST_HW_VERSION and falsely flag
    every VM in that workbook as running an outdated hardware version. This
    normalizes on the numeric value itself so a stray blank cell elsewhere
    in the column can't corrupt the rest of the values.
    """
    raw = normalize_strings(series, default="unknown")
    numeric = pd.to_numeric(series, errors="coerce")
    whole_number = numeric.notna() & numeric.eq(numeric.round())
    normalized = raw.copy()
    normalized.loc[whole_number] = numeric.loc[whole_number].astype("int64").astype(str)
    return normalized


def get_normalized_hw_version_column(frame: pd.DataFrame, column: str = "HW version") -> pd.Series:
    if column not in frame.columns:
        return pd.Series(["unknown"] * len(frame.index), index=frame.index, dtype="object")
    return normalize_hw_version_series(frame[column])


def split_multi_value_cell(value: object) -> list[str]:
    if pd.isna(value):
        return []
    parts = [
        item.strip()
        for item in re.split(r"[,;\n|]+", str(value))
    ]
    return [item for item in parts if item]


def extract_ipv4_addresses(value: object) -> list[str]:
    if pd.isna(value):
        return []
    return re.findall(r"\b(?:\d{1,3}\.){3}\d{1,3}\b", str(value))


def is_excluded_duplicate_ip(ip_address: str) -> bool:
    return bool(re.fullmatch(r"169\.254\.\d{1,3}\.\d{1,3}", ip_address.strip()))


def numeric_series(frame: pd.DataFrame, column: str) -> pd.Series:
    if column not in frame.columns:
        return pd.Series([0] * len(frame.index), index=frame.index, dtype="float64")
    return pd.to_numeric(frame[column], errors="coerce").fillna(0)


def min_positive_or_zero(series: pd.Series) -> float:
    positive_values = pd.to_numeric(series, errors="coerce").dropna()
    positive_values = positive_values.loc[positive_values.gt(0)]
    if positive_values.empty:
        return 0.0
    return float(positive_values.min())


def datetime_series(frame: pd.DataFrame, column: str) -> pd.Series:
    if column not in frame.columns:
        return pd.Series([pd.NaT] * len(frame.index), index=frame.index, dtype="datetime64[ns]")
    return pd.to_datetime(frame[column], errors="coerce")


def get_normalized_column(
    frame: pd.DataFrame,
    column: str,
    *,
    default: str = "",
    bool_values: bool = False,
) -> pd.Series:
    if column not in frame.columns:
        return pd.Series([default] * len(frame.index), index=frame.index, dtype="object")
    if bool_values:
        return normalize_bool_strings(frame[column])
    return normalize_strings(frame[column], default=default)


def datastore_non_local_mask(frame: pd.DataFrame) -> pd.Series:
    """Return True for datastores considered shared/non-local.

    Prefers RVTools' own 'MHA' (multiple-host access) and '# Hosts' signals.
    If an export has neither column, falls back to a name-based guess (only
    the default ESXi 'datastore1' name is treated as local) and logs a
    warning, since that guess can misclassify any other locally-named
    datastore as shared and skew ds_cross_clustered/ds_with_no_vms and the
    per-cluster shared/local datastore counts for that workbook.
    """
    if frame.empty:
        return pd.Series(dtype="bool")

    has_mha = "MHA" in frame.columns
    has_hosts = "# Hosts" in frame.columns

    if has_mha:
        mha_true = get_normalized_column(frame, "MHA", default="blank", bool_values=True).eq("true")
    else:
        mha_true = pd.Series([False] * len(frame.index), index=frame.index, dtype="bool")

    if has_hosts:
        multi_host = numeric_series(frame, "# Hosts").gt(1)
    else:
        multi_host = pd.Series([False] * len(frame.index), index=frame.index, dtype="bool")

    if has_mha and has_hosts:
        return mha_true | multi_host
    if has_mha:
        return mha_true
    if has_hosts:
        return multi_host

    # Last-resort fallback for exports without locality signals.
    logging.getLogger("rvtools_summary").warning(
        "vDatastore sheet has neither 'MHA' nor '# Hosts' column; falling back to "
        "a name-based guess for datastore locality (only a 'datastore1' name is "
        "treated as local). Datastore locality/cluster metrics may be inaccurate "
        "for this workbook.",
        extra={"console": False},
    )
    datastore_names = get_normalized_column(frame, "Name", default="")
    return ~datastore_names.str.lower().str.fullmatch(r"datastore1(?:\s+\(\d+\))?")


def datastore_vsan_mask(frame: pd.DataFrame) -> pd.Series:
    if frame.empty:
        return pd.Series(dtype="bool")

    # RVTools exposes datastore type directly on vDatastore; use that as the
    # authoritative signal for vSAN classification.
    datastore_type = get_normalized_column(frame, "Type", default="").str.lower()
    return datastore_type.str.contains("vsan", regex=False)


def normalize_column_name(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "_", value.strip().lower()).strip("_")
    return normalized or "unknown"


def controller_count_columns(counter: Counter[str], controller_keys: list[str]) -> dict[str, int]:
    return {
        f"vms_scsi_controller_{normalize_column_name(key)}_count": int(counter.get(key, 0))
        for key in controller_keys
    }


def adapter_count_columns(counter: Counter[str], adapter_keys: list[str]) -> dict[str, int]:
    return {
        f"vms_network_adapter_{normalize_column_name(key)}_count": int(counter.get(key, 0))
        for key in adapter_keys
    }


def excel_sheet_name(sheet_name: str) -> str:
    return EXCEL_SHEET_NAME_ALIASES.get(sheet_name, sheet_name)


def add_summary_header_links(workbook, summary_columns: list[str]) -> None:
    if not OPENPYXL_AVAILABLE or Font is None or get_column_letter is None:
        return
    summary_sheet = workbook["Summary"]
    hyperlink_font = Font(color="0563C1", underline="single", bold=True)
    for index, column_name in enumerate(summary_columns, start=1):
        target_sheet = SUMMARY_HEADER_SHEET_LINKS.get(column_name)
        if not target_sheet:
            continue
        target_sheet = excel_sheet_name(target_sheet)
        cell = summary_sheet.cell(row=1, column=index)
        cell.hyperlink = f"#'{target_sheet}'!A1"
        cell.font = hyperlink_font


def add_detail_sheet_back_links(workbook, summary_columns: list[str], sheet_names: Iterable[str]) -> None:
    if not OPENPYXL_AVAILABLE or Font is None or get_column_letter is None:
        return
    hyperlink_font = Font(color="0563C1", underline="single", bold=True)
    summary_columns_by_sheet: dict[str, list[tuple[str, int]]] = {}
    for index, column_name in enumerate(summary_columns, start=1):
        target_sheet = SUMMARY_HEADER_SHEET_LINKS.get(column_name)
        if not target_sheet:
            continue
        summary_columns_by_sheet.setdefault(excel_sheet_name(target_sheet), []).append((column_name, index))

    for sheet_name in sheet_names:
        worksheet = workbook[excel_sheet_name(sheet_name)]
        summary_links = summary_columns_by_sheet.get(worksheet.title, [])
        if not summary_links:
            continue
        for offset, (column_name, summary_column_index) in enumerate(summary_links, start=1):
            back_cell = worksheet.cell(row=1, column=worksheet.max_column + offset)
            back_cell.value = (
                "back_to_summary"
                if len(summary_links) == 1
                else f"back_to_{column_name}"
            )
            summary_column_letter = get_column_letter(summary_column_index)
            back_cell.hyperlink = f"#'Summary'!{summary_column_letter}1"
            back_cell.font = hyperlink_font


def build_counter_breakdown_sheet(
    summaries: list[WorkbookSummary],
    combined_summary: WorkbookSummary | None,
    counter_attr: str,
) -> pd.DataFrame:
    values = sorted(
        {
            value
            for summary in summaries
            for value in getattr(summary, counter_attr)
        }
        | (
            set(getattr(combined_summary, counter_attr))
            if combined_summary is not None
            else set()
        )
    )
    rows: list[dict[str, object]] = []
    for summary in summaries:
        row = {"source_vcenter": summary.source_vcenter}
        counter = getattr(summary, counter_attr)
        for value in values:
            row[value] = int(counter.get(value, 0))
        rows.append(row)

    grand_total = {"source_vcenter": "grand_total"}
    total_counts = (
        getattr(combined_summary, counter_attr)
        if combined_summary is not None
        else Counter()
    )
    for value in values:
        grand_total[value] = int(total_counts.get(value, 0))
    rows.append(grand_total)

    if not values:
        return pd.DataFrame(columns=["source_vcenter"])
    return pd.DataFrame(rows, columns=["source_vcenter", *values])


def build_host_to_cluster_map(vhost: pd.DataFrame) -> dict[str, str]:
    if {"Host", "Cluster"}.issubset(vhost.columns):
        return (
            pd.DataFrame(
                {
                    "host_name": get_normalized_column(vhost, "Host", default=""),
                    "cluster_name": get_normalized_column(vhost, "Cluster", default=""),
                }
            )
            .loc[lambda df: df["host_name"].ne("") & df["cluster_name"].ne("")]
            .drop_duplicates(subset=["host_name"], keep="first")
            .set_index("host_name")["cluster_name"]
            .to_dict()
        )
    return {}


def build_resolved_datastore_cluster_frame(
    vdatastore: pd.DataFrame,
    host_to_cluster: dict[str, str],
) -> pd.DataFrame:
    columns = [
        "source_vcenter",
        "datastore_name",
        "datastore_identifier",
        "MHA",
        "Hosts Reported",
        "Hosts",
        "Reported Cluster Name",
        "resolved_cluster_list",
        "Cluster Names",
        "Cluster Count",
        "Datastore Locality",
        "Capacity MiB",
        "is_non_local",
        "is_vsan",
    ]
    if "Name" not in vdatastore.columns:
        return pd.DataFrame(columns=columns)

    datastore_identifier_column = "Object ID" if "Object ID" in vdatastore.columns else "Name"
    datastore_frame = pd.DataFrame(
        {
            "source_vcenter": get_normalized_column(vdatastore, "VI SDK Server", default="unknown"),
            "datastore_identifier": get_normalized_column(
                vdatastore,
                datastore_identifier_column,
                default="",
            ),
            "datastore_name": get_normalized_column(vdatastore, "Name", default=""),
            "MHA": get_normalized_column(vdatastore, "MHA", default="blank", bool_values=True),
            "Hosts Reported": numeric_series(vdatastore, "# Hosts"),
            "Hosts": get_normalized_column(vdatastore, "Hosts", default=""),
            "Reported Cluster Name": get_normalized_column(vdatastore, "Cluster name", default=""),
            "Capacity MiB": numeric_series(vdatastore, "Capacity MiB"),
            "is_vsan": datastore_vsan_mask(vdatastore),
        }
    ).loc[lambda df: df["datastore_name"].ne("")]
    if datastore_frame.empty:
        return pd.DataFrame(columns=columns)

    datastore_frame["datastore_identifier"] = datastore_frame["datastore_identifier"].mask(
        datastore_frame["datastore_identifier"].eq(""),
        datastore_frame["datastore_name"],
    )
    non_local_mask = datastore_non_local_mask(vdatastore).reindex(
        datastore_frame.index,
        fill_value=False,
    )

    aggregated: dict[tuple[str, str], dict[str, object]] = {}
    for record, is_non_local in zip(datastore_frame.to_dict("records"), non_local_mask.tolist()):
        key = (str(record["source_vcenter"]), str(record["datastore_identifier"]))
        entry = aggregated.setdefault(
            key,
            {
                "source_vcenter": record["source_vcenter"],
                "datastore_name": record["datastore_name"],
                "datastore_identifier": record["datastore_identifier"],
                "mha_values": set(),
                "host_names": set(),
                "resolved_clusters": set(),
                "reported_cluster_names": set(),
                "reported_hosts_max": 0,
                "capacity_mib_max": 0.0,
                "is_non_local": False,
                "is_vsan": False,
            },
        )
        entry["datastore_name"] = record["datastore_name"]
        entry["reported_hosts_max"] = max(entry["reported_hosts_max"], int(record["Hosts Reported"]))
        entry["capacity_mib_max"] = max(entry["capacity_mib_max"], float(record["Capacity MiB"]))
        entry["is_non_local"] = bool(entry["is_non_local"] or is_non_local)
        entry["is_vsan"] = bool(entry["is_vsan"] or record["is_vsan"])
        if record["MHA"] != "blank":
            entry["mha_values"].add(record["MHA"])
        if record["Reported Cluster Name"]:
            entry["reported_cluster_names"].update(split_multi_value_cell(record["Reported Cluster Name"]))

        for host_name in split_multi_value_cell(record["Hosts"]):
            entry["host_names"].add(host_name)
            cluster_name = host_to_cluster.get(host_name)
            if cluster_name:
                entry["resolved_clusters"].add(cluster_name)

    rows: list[dict[str, object]] = []
    for entry in aggregated.values():
        resolved_cluster_list = sorted(entry["resolved_clusters"] or entry["reported_cluster_names"])
        host_names = sorted(entry["host_names"])
        rows.append(
            {
                "source_vcenter": entry["source_vcenter"],
                "datastore_name": entry["datastore_name"],
                "datastore_identifier": entry["datastore_identifier"],
                "MHA": ", ".join(sorted(entry["mha_values"])) if entry["mha_values"] else "blank",
                "Hosts Reported": int(entry["reported_hosts_max"]),
                "Hosts": ", ".join(host_names),
                "Reported Cluster Name": ", ".join(sorted(entry["reported_cluster_names"])),
                "resolved_cluster_list": resolved_cluster_list,
                "Cluster Names": ", ".join(resolved_cluster_list),
                "Cluster Count": len(resolved_cluster_list),
                "Datastore Locality": "shared/non-local" if entry["is_non_local"] else "local/non-shared",
                "Capacity MiB": float(entry["capacity_mib_max"]),
                "is_non_local": bool(entry["is_non_local"]),
                "is_vsan": bool(entry["is_vsan"]),
            }
        )

    return pd.DataFrame(rows, columns=columns).sort_values(
        by=["source_vcenter", "datastore_name", "datastore_identifier"],
        kind="stable",
    ).reset_index(drop=True)


def build_datastore_inventory_frame(vdatastore: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "source_vcenter",
        "datastore_name",
        "datastore_identifier",
        "is_non_local",
        "Type",
        "Accessible",
        "# VMs total",
        "# Hosts",
        "Hosts",
        "Cluster name",
        "Capacity MiB",
        "Provisioned MiB",
        "In Use MiB",
        "Free MiB",
        "Free %",
        "# Extents",
        "URL",
        "Object ID",
        "is_vsan",
    ]
    if "Name" not in vdatastore.columns:
        return pd.DataFrame(columns=columns)

    datastore_identifier_column = "Object ID" if "Object ID" in vdatastore.columns else "Name"
    non_local_mask = datastore_non_local_mask(vdatastore)
    frame = pd.DataFrame(
        {
            "source_vcenter": get_normalized_column(
                vdatastore,
                "VI SDK Server",
                default="unknown",
            ),
            "datastore_name": get_normalized_column(vdatastore, "Name", default=""),
            "datastore_identifier": get_normalized_column(
                vdatastore,
                datastore_identifier_column,
                default="",
            ),
            "is_non_local": non_local_mask.reindex(vdatastore.index, fill_value=False).astype(bool),
            "Type": get_normalized_column(vdatastore, "Type", default="unknown"),
            "Accessible": get_normalized_column(
                vdatastore,
                "Accessible",
                default="blank",
                bool_values=True,
            ),
            "# VMs total": numeric_series(vdatastore, "# VMs total"),
            "# Hosts": numeric_series(vdatastore, "# Hosts"),
            "Hosts": get_normalized_column(vdatastore, "Hosts", default=""),
            "Cluster name": get_normalized_column(vdatastore, "Cluster name", default=""),
            "Capacity MiB": numeric_series(vdatastore, "Capacity MiB"),
            "Provisioned MiB": numeric_series(vdatastore, "Provisioned MiB"),
            "In Use MiB": numeric_series(vdatastore, "In Use MiB"),
            "Free MiB": numeric_series(vdatastore, "Free MiB"),
            "Free %": numeric_series(vdatastore, "Free %"),
            "# Extents": numeric_series(vdatastore, "# Extents"),
            "URL": get_normalized_column(vdatastore, "URL", default=""),
            "Object ID": get_normalized_column(vdatastore, "Object ID", default=""),
            "is_vsan": datastore_vsan_mask(vdatastore),
        }
    ).loc[lambda df: df["datastore_name"].ne("")]
    if frame.empty:
        return pd.DataFrame(columns=columns)

    frame["datastore_identifier"] = frame["datastore_identifier"].mask(
        frame["datastore_identifier"].eq(""),
        frame["datastore_name"],
    )
    return (
        frame.drop_duplicates(
            subset=["source_vcenter", "datastore_identifier"],
            keep="first",
        )
        .sort_values(by=["source_vcenter", "datastore_name"], kind="stable")
        .reset_index(drop=True)
    )


def datastore_report_columns() -> list[str]:
    return [
        "source_vcenter",
        "datastore_name",
        "Type",
        "Accessible",
        "# VMs total",
        "# Hosts",
        "Hosts",
        "Cluster name",
        "Capacity MiB",
        "Provisioned MiB",
        "In Use MiB",
        "Free MiB",
        "Free %",
        "# Extents",
        "URL",
        "Object ID",
    ]


def build_host_state_map(vhost: pd.DataFrame | None) -> dict[str, str]:
    if vhost is None or "Host" not in vhost.columns:
        return {}

    host_frame = pd.DataFrame(
        {
            "host_name": get_normalized_column(vhost, "Host", default=""),
        }
    ).loc[lambda df: df["host_name"].ne("")]
    if host_frame.empty:
        return {}

    if "Connection state" in vhost.columns:
        host_frame["connection_state"] = get_normalized_column(vhost, "Connection state", default="unknown")
    else:
        host_frame["connection_state"] = "unknown"

    if "in Maintenance Mode" in vhost.columns:
        host_frame["maintenance_mode"] = get_normalized_column(
            vhost,
            "in Maintenance Mode",
            default="blank",
            bool_values=True,
        )
    else:
        host_frame["maintenance_mode"] = "blank"

    def host_state_for(row: pd.Series) -> str:
        if row["maintenance_mode"] == "true":
            return "maintenanceMode"
        if row["connection_state"] not in {"", "unknown"}:
            return str(row["connection_state"])
        if row["maintenance_mode"] == "false":
            return "notInMaintenanceMode"
        return "unknown"

    host_frame["state"] = host_frame.apply(host_state_for, axis=1)
    return (
        host_frame.drop_duplicates(subset=["host_name"], keep="first")
        .set_index("host_name")["state"]
        .to_dict()
    )


def build_cluster_summary_report(workbooks: list[Path]) -> pd.DataFrame:
    rows: list[dict[str, object]] = []

    for workbook_path in workbooks:
        workbook = pd.ExcelFile(workbook_path)

        cluster_rows: list[dict[str, object]] = []
        host_to_cluster: dict[str, str] = {}
        source_vcenter = "unknown"

        if "vHost" in workbook.sheet_names:
            vhost = load_sheet(workbook_path, "vHost")
            host_cluster_frame = pd.DataFrame(
                {
                    "source_vcenter": get_normalized_column(vhost, "VI SDK Server", default="unknown"),
                    "host_name": get_normalized_column(vhost, "Host", default=""),
                    "cluster_name": get_normalized_column(vhost, "Cluster", default=""),
                }
            ).loc[lambda df: df["host_name"].ne("") & df["cluster_name"].ne("")]

            if not host_cluster_frame.empty:
                source_vcenter = host_cluster_frame["source_vcenter"].iloc[0]
                host_to_cluster = build_host_to_cluster_map(vhost)
                cluster_rows.extend(
                    host_cluster_frame.groupby("cluster_name", dropna=False)["host_name"]
                    .nunique()
                    .reset_index(name="hosts_total")
                    .assign(
                        source_vcenter=source_vcenter,
                        shared_datastores=0,
                        vsan_datastores=0,
                        local_non_shared_datastores=0,
                        total_datastores=0,
                        vms_total=0,
                        vms_powered_on=0,
                        vms_powered_off=0,
                        max_shared_datastore_size=0.0,
                        min_shared_datastore_size=0.0,
                        max_vsan_datastore_size=0.0,
                        min_vsan_datastore_size=0.0,
                        vms_max_size=0.0,
                        vms_min_size=0.0,
                        vms_max_storage_used=0.0,
                        vms_min_storage_used=0.0,
                    )
                    .rename(columns={"cluster_name": "cluster"})
                    .to_dict("records")
                )

        if not cluster_rows and "vCluster" in workbook.sheet_names:
            vcluster = load_sheet(workbook_path, "vCluster")
            vcluster_frame = pd.DataFrame(
                {
                    "source_vcenter": get_normalized_column(vcluster, "VI SDK Server", default="unknown"),
                    "cluster": get_normalized_column(vcluster, "Name", default=""),
                    "hosts_total": numeric_series(vcluster, "NumHosts"),
                }
            ).loc[lambda df: df["cluster"].ne("")]
            if not vcluster_frame.empty:
                source_vcenter = vcluster_frame["source_vcenter"].iloc[0]
                cluster_rows.extend(
                    vcluster_frame.assign(
                        shared_datastores=0,
                        vsan_datastores=0,
                        local_non_shared_datastores=0,
                        total_datastores=0,
                        vms_total=0,
                        vms_powered_on=0,
                        vms_powered_off=0,
                        max_shared_datastore_size=0.0,
                        min_shared_datastore_size=0.0,
                        max_vsan_datastore_size=0.0,
                        min_vsan_datastore_size=0.0,
                        vms_max_size=0.0,
                        vms_min_size=0.0,
                        vms_max_storage_used=0.0,
                        vms_min_storage_used=0.0,
                    ).to_dict("records")
                )

        cluster_frame = pd.DataFrame(cluster_rows)
        if cluster_frame.empty:
            continue

        cluster_frame = (
            cluster_frame.groupby(["source_vcenter", "cluster"], dropna=False, as_index=False)
            .agg(
                hosts_total=("hosts_total", "max"),
                shared_datastores=("shared_datastores", "max"),
                vsan_datastores=("vsan_datastores", "max"),
                local_non_shared_datastores=("local_non_shared_datastores", "max"),
                total_datastores=("total_datastores", "max"),
                vms_total=("vms_total", "max"),
                vms_powered_on=("vms_powered_on", "max"),
                vms_powered_off=("vms_powered_off", "max"),
                max_shared_datastore_size=("max_shared_datastore_size", "max"),
                min_shared_datastore_size=("min_shared_datastore_size", "max"),
                max_vsan_datastore_size=("max_vsan_datastore_size", "max"),
                min_vsan_datastore_size=("min_vsan_datastore_size", "max"),
                vms_max_size=("vms_max_size", "max"),
                vms_min_size=("vms_min_size", "max"),
                vms_max_storage_used=("vms_max_storage_used", "max"),
                vms_min_storage_used=("vms_min_storage_used", "max"),
            )
        )

        if "vInfo" in workbook.sheet_names:
            vinfo = load_sheet(workbook_path, "vInfo")
            vm_cluster_frame = pd.DataFrame(
                {
                    "source_vcenter": get_normalized_column(vinfo, "VI SDK Server", default=source_vcenter),
                    "cluster": get_normalized_column(vinfo, "Cluster", default=""),
                    "vm_id": get_normalized_column(vinfo, "VM ID", default=""),
                    "powerstate": get_normalized_column(vinfo, "Powerstate", default="unknown"),
                    "vms_size_mib": numeric_series(vinfo, "Provisioned MiB"),
                    "vms_storage_used_mib": numeric_series(vinfo, "In Use MiB"),
                }
            ).loc[lambda df: df["cluster"].ne("")]
            if not vm_cluster_frame.empty:
                vm_cluster_frame["vm_identifier"] = vm_cluster_frame["vm_id"].mask(
                    vm_cluster_frame["vm_id"].eq(""),
                    vm_cluster_frame.index.astype(str),
                )
                vm_cluster_frame = (
                    vm_cluster_frame.sort_values(
                        by=["source_vcenter", "cluster", "vm_identifier", "vms_size_mib"],
                        ascending=[True, True, True, False],
                        kind="stable",
                    )
                    .drop_duplicates(
                        subset=["source_vcenter", "cluster", "vm_identifier"],
                        keep="first",
                    )
                )
                vm_cluster_counts = (
                    vm_cluster_frame.groupby(["source_vcenter", "cluster"], as_index=False)
                    .agg(
                        vms_total=("vm_identifier", "nunique"),
                        vms_powered_on=("powerstate", lambda s: int(s.eq("poweredOn").sum())),
                        vms_powered_off=("powerstate", lambda s: int(s.eq("poweredOff").sum())),
                        vms_max_size=("vms_size_mib", "max"),
                        vms_min_size=("vms_size_mib", min_positive_or_zero),
                        vms_max_storage_used=("vms_storage_used_mib", "max"),
                        vms_min_storage_used=("vms_storage_used_mib", min_positive_or_zero),
                    )
                )
                cluster_frame = cluster_frame.merge(
                    vm_cluster_counts,
                    on=["source_vcenter", "cluster"],
                    how="left",
                    suffixes=("", "_computed"),
                )
                for column_name in [
                    "vms_total",
                    "vms_powered_on",
                    "vms_powered_off",
                    "vms_max_size",
                    "vms_min_size",
                    "vms_max_storage_used",
                    "vms_min_storage_used",
                ]:
                    cluster_frame[column_name] = cluster_frame[f"{column_name}_computed"].combine_first(
                        cluster_frame[column_name]
                    )
                    cluster_frame = cluster_frame.drop(columns=[f"{column_name}_computed"])

        if "vDatastore" in workbook.sheet_names:
            vdatastore = load_sheet(workbook_path, "vDatastore")
            resolved_datastore_frame = build_resolved_datastore_cluster_frame(vdatastore, host_to_cluster)
            if not resolved_datastore_frame.empty:
                shared_datastore_pairs = {
                    (cluster_name, record["datastore_identifier"])
                    for record in resolved_datastore_frame.loc[
                        resolved_datastore_frame["is_non_local"]
                    ].to_dict("records")
                    for cluster_name in record["resolved_cluster_list"]
                }
                if shared_datastore_pairs:
                    shared_datastore_counts = (
                        pd.DataFrame(
                            shared_datastore_pairs,
                            columns=["cluster", "datastore_identifier"],
                        )
                        .groupby("cluster", as_index=False)["datastore_identifier"]
                        .nunique()
                        .rename(columns={"datastore_identifier": "shared_datastores"})
                    )
                    cluster_frame = cluster_frame.merge(
                        shared_datastore_counts,
                        on="cluster",
                        how="left",
                        suffixes=("", "_computed"),
                    )
                    cluster_frame["shared_datastores"] = (
                        cluster_frame["shared_datastores_computed"].fillna(cluster_frame["shared_datastores"])
                    )
                    cluster_frame = cluster_frame.drop(columns=["shared_datastores_computed"])

                    shared_datastore_sizes = pd.DataFrame(
                        [
                            {
                                "cluster": cluster_name,
                                "datastore_identifier": record["datastore_identifier"],
                                "datastore_size_mib": float(record["Capacity MiB"]),
                            }
                            for record in resolved_datastore_frame.loc[
                                resolved_datastore_frame["is_non_local"]
                            ].to_dict("records")
                            for cluster_name in record["resolved_cluster_list"]
                        ]
                    )
                    if not shared_datastore_sizes.empty:
                        shared_datastore_sizes = shared_datastore_sizes.drop_duplicates(
                            subset=["cluster", "datastore_identifier"],
                            keep="first",
                        )
                        shared_datastore_size_summary = (
                            shared_datastore_sizes.groupby("cluster", as_index=False)
                            .agg(
                                max_shared_datastore_size=("datastore_size_mib", "max"),
                                min_shared_datastore_size=("datastore_size_mib", "min"),
                            )
                        )
                        cluster_frame = cluster_frame.merge(
                            shared_datastore_size_summary,
                            on="cluster",
                            how="left",
                            suffixes=("", "_computed"),
                        )
                        for column_name in ["max_shared_datastore_size", "min_shared_datastore_size"]:
                            cluster_frame[column_name] = cluster_frame[f"{column_name}_computed"].fillna(
                                cluster_frame[column_name]
                            )
                            cluster_frame = cluster_frame.drop(columns=[f"{column_name}_computed"])

                vsan_datastore_records = [
                    {
                        "cluster": cluster_name,
                        "datastore_identifier": record["datastore_identifier"],
                        "datastore_name": record["datastore_name"],
                        "datastore_size_mib": float(record["Capacity MiB"]),
                    }
                    for record in resolved_datastore_frame.loc[
                        resolved_datastore_frame["is_vsan"]
                    ].to_dict("records")
                    for cluster_name in record["resolved_cluster_list"]
                ]
                if vsan_datastore_records:
                    vsan_datastore_frame = pd.DataFrame(vsan_datastore_records).drop_duplicates(
                        subset=["cluster", "datastore_identifier"],
                        keep="first",
                    )
                    vsan_datastore_counts = (
                        vsan_datastore_frame.groupby("cluster", as_index=False)["datastore_identifier"]
                        .nunique()
                        .rename(columns={"datastore_identifier": "vsan_datastores"})
                    )
                    cluster_frame = cluster_frame.merge(
                        vsan_datastore_counts,
                        on="cluster",
                        how="left",
                        suffixes=("", "_computed"),
                    )
                    cluster_frame["vsan_datastores"] = (
                        cluster_frame["vsan_datastores_computed"].fillna(cluster_frame["vsan_datastores"])
                    )
                    cluster_frame = cluster_frame.drop(columns=["vsan_datastores_computed"])

                    vsan_datastore_size_summary = (
                        vsan_datastore_frame.groupby("cluster", as_index=False)
                        .agg(
                            max_vsan_datastore_size=("datastore_size_mib", "max"),
                            min_vsan_datastore_size=("datastore_size_mib", "min"),
                        )
                    )
                    cluster_frame = cluster_frame.merge(
                        vsan_datastore_size_summary,
                        on="cluster",
                        how="left",
                        suffixes=("", "_computed"),
                    )
                    for column_name in ["max_vsan_datastore_size", "min_vsan_datastore_size"]:
                        cluster_frame[column_name] = cluster_frame[f"{column_name}_computed"].fillna(
                            cluster_frame[column_name]
                        )
                        cluster_frame = cluster_frame.drop(columns=[f"{column_name}_computed"])

                local_datastore_pairs = {
                    (cluster_name, record["datastore_identifier"])
                    for record in resolved_datastore_frame.loc[
                        ~resolved_datastore_frame["is_non_local"]
                    ].to_dict("records")
                    for cluster_name in record["resolved_cluster_list"]
                }
                if local_datastore_pairs:
                    local_datastore_counts = (
                        pd.DataFrame(
                            local_datastore_pairs,
                            columns=["cluster", "datastore_identifier"],
                        )
                        .groupby("cluster", as_index=False)["datastore_identifier"]
                        .nunique()
                        .rename(
                            columns={
                                "datastore_identifier": "local_non_shared_datastores"
                            }
                        )
                    )
                    cluster_frame = cluster_frame.merge(
                        local_datastore_counts,
                        on="cluster",
                        how="left",
                        suffixes=("", "_computed"),
                    )
                    cluster_frame["local_non_shared_datastores"] = (
                        cluster_frame["local_non_shared_datastores_computed"].fillna(
                            cluster_frame["local_non_shared_datastores"]
                        )
                    )
                    cluster_frame = cluster_frame.drop(
                        columns=["local_non_shared_datastores_computed"]
                    )

        rows.extend(
            cluster_frame.assign(shared_datastores=lambda df: df["shared_datastores"].fillna(0).astype(int))
            .assign(vsan_datastores=lambda df: df["vsan_datastores"].fillna(0).astype(int))
            .assign(
                local_non_shared_datastores=lambda df: df["local_non_shared_datastores"]
                .fillna(0)
                .astype(int)
            )
            .assign(
                total_datastores=lambda df: (
                    df["shared_datastores"] + df["local_non_shared_datastores"]
                ).astype(int)
            )
            .assign(hosts_total=lambda df: df["hosts_total"].fillna(0).astype(int))
            .assign(vms_total=lambda df: df["vms_total"].fillna(0).astype(int))
            .assign(vms_powered_on=lambda df: df["vms_powered_on"].fillna(0).astype(int))
            .assign(vms_powered_off=lambda df: df["vms_powered_off"].fillna(0).astype(int))
            .assign(
                max_shared_datastore_size=lambda df: df["max_shared_datastore_size"].fillna(0.0)
            )
            .assign(
                min_shared_datastore_size=lambda df: df["min_shared_datastore_size"].fillna(0.0)
            )
            .assign(
                max_vsan_datastore_size=lambda df: df["max_vsan_datastore_size"].fillna(0.0)
            )
            .assign(
                min_vsan_datastore_size=lambda df: df["min_vsan_datastore_size"].fillna(0.0)
            )
            .assign(vms_max_size=lambda df: df["vms_max_size"].fillna(0.0))
            .assign(vms_min_size=lambda df: df["vms_min_size"].fillna(0.0))
            .assign(vms_max_storage_used=lambda df: df["vms_max_storage_used"].fillna(0.0))
            .assign(vms_min_storage_used=lambda df: df["vms_min_storage_used"].fillna(0.0))
            .assign(
                **{
                    "max_shared_datastore_size (GB)": lambda df: (
                        df["max_shared_datastore_size"] / MIB_PER_GIB
                    ).round(2),
                    "min_shared_datastore_size (GB)": lambda df: (
                        df["min_shared_datastore_size"] / MIB_PER_GIB
                    ).round(2),
                    "max_vsan_datastore_size (GB)": lambda df: (
                        df["max_vsan_datastore_size"] / MIB_PER_GIB
                    ).round(2),
                    "min_vsan_datastore_size (GB)": lambda df: (
                        df["min_vsan_datastore_size"] / MIB_PER_GIB
                    ).round(2),
                    "vms_max_size (GB)": lambda df: (df["vms_max_size"] / MIB_PER_GIB).round(2),
                    "vms_min_size (GB)": lambda df: (df["vms_min_size"] / MIB_PER_GIB).round(2),
                    "vms_max_storage_used (GB)": lambda df: (
                        df["vms_max_storage_used"] / MIB_PER_GIB
                    ).round(2),
                    "vms_min_storage_used (GB)": lambda df: (
                        df["vms_min_storage_used"] / MIB_PER_GIB
                    ).round(2),
                }
            )
            .drop(
                columns=[
                    "max_shared_datastore_size",
                    "min_shared_datastore_size",
                    "max_vsan_datastore_size",
                    "min_vsan_datastore_size",
                    "vms_max_size",
                    "vms_min_size",
                    "vms_max_storage_used",
                    "vms_min_storage_used",
                ]
            )
            .sort_values(by=["source_vcenter", "cluster"], kind="stable")
            .to_dict("records")
        )

    return pd.DataFrame(
        rows,
        columns=[
            "source_vcenter",
            "cluster",
            "hosts_total",
            "total_datastores",
            "shared_datastores",
            "vsan_datastores",
            "local_non_shared_datastores",
            "max_shared_datastore_size (GB)",
            "min_shared_datastore_size (GB)",
            "max_vsan_datastore_size (GB)",
            "min_vsan_datastore_size (GB)",
            "vms_total",
            "vms_powered_on",
            "vms_powered_off",
            "vms_max_size (GB)",
            "vms_min_size (GB)",
            "vms_max_storage_used (GB)",
            "vms_min_storage_used (GB)",
        ],
    )


@dataclass
class AviNetworkSummary:
    """Avi/NSX ALB network entities pulled from --avi-data: SE data-plane IPs
    and VsVip VIP addresses, shaped both as Counters (for the same global
    duplicate-IP reconciliation used for vSphere hosts/VMs) and as rich
    per-entity records (for the Avi detail sheets and for folding into the
    RVTools-side duplicate_ips sheet, so an Avi VIP colliding with a vSphere
    VM's IP shows up in both places consistently).

    MAC addresses are intentionally not tracked here: an SE's management IP
    and every MAC it reports (mgmt + data) belong to a VM RVTools already
    inventories on its own, so treating them as a separate Avi entity would
    just create a false duplicate against that same VM.
    """

    raw_reports: dict[str, pd.DataFrame] = field(default_factory=dict)
    service_engines_total: int = 0
    vsvips_total: int = 0
    virtual_services_total: int = 0
    network_subnets_total: int = 0
    se_ip_counts: Counter[str] = field(default_factory=Counter)
    vip_counts: Counter[str] = field(default_factory=Counter)
    se_ip_records: list[dict[str, object]] = field(default_factory=list)
    vip_records: list[dict[str, object]] = field(default_factory=list)
    combined_ip_records: list[dict[str, object]] = field(default_factory=list)
    se_with_duplicate_ip: int = 0
    vip_with_duplicate_ip: int = 0


def build_avi_network_summary(avi_reports: dict[str, pd.DataFrame]) -> AviNetworkSummary:
    avi = AviNetworkSummary(raw_reports=avi_reports)

    service_engines = avi_reports.get("service_engines")
    if service_engines is not None and {"uuid", "se_name"}.issubset(service_engines.columns):
        avi.service_engines_total = int(len(service_engines.index))
        se_frame = pd.DataFrame(
            {
                "controller": get_normalized_column(service_engines, "controller", default="unknown"),
                "se_name": get_normalized_column(service_engines, "se_name", default=""),
                "uuid": get_normalized_column(service_engines, "uuid", default=""),
                "se_group": get_normalized_column(service_engines, "se_group", default=""),
                "data_ips": get_normalized_column(service_engines, "data_ips", default=""),
            }
        ).loc[lambda df: df["uuid"].ne("") | df["se_name"].ne("")]

        # Only the SE's data-plane IPs are Avi-specific. Its management IP
        # and all of its MAC addresses (mgmt + data) belong to a VM that
        # RVTools already inventories directly, so folding those in here
        # too would double-count the same NIC as two different "entities"
        # (once as a vSphere VM, once as an Avi SE) and produce a false
        # duplicate between them.
        for record in se_frame.to_dict("records"):
            entity_id = record["uuid"] or record["se_name"]
            ip_values = set(split_multi_value_cell(record["data_ips"]))
            for ip_address in ip_values:
                avi.se_ip_records.append(
                    {
                        "controller": record["controller"],
                        "se_name": record["se_name"],
                        "uuid": record["uuid"],
                        "se_group": record["se_group"],
                        "IP Address": ip_address,
                    }
                )
                avi.combined_ip_records.append(
                    {
                        "source": record["controller"],
                        "name": record["se_name"],
                        "state": "avi_service_engine",
                        "ip_address": ip_address,
                        "mac_address": "",
                        "id": entity_id,
                        "network/port_group": record["se_group"],
                        "nic_label/device": "avi_se",
                    }
                )
            avi.se_ip_counts.update(ip_values)

    vsvips = avi_reports.get("vsvips")
    if vsvips is not None and {"uuid", "vsvip_name", "vip_addresses"}.issubset(vsvips.columns):
        avi.vsvips_total = int(len(vsvips.index))
        vsvip_frame = pd.DataFrame(
            {
                "controller": get_normalized_column(vsvips, "controller", default="unknown"),
                "vsvip_name": get_normalized_column(vsvips, "vsvip_name", default=""),
                "uuid": get_normalized_column(vsvips, "uuid", default=""),
                "vip_addresses": get_normalized_column(vsvips, "vip_addresses", default=""),
                "vip_networks": get_normalized_column(vsvips, "vip_networks", default=""),
            }
        ).loc[lambda df: df["uuid"].ne("") | df["vsvip_name"].ne("")]

        for record in vsvip_frame.to_dict("records"):
            entity_id = record["uuid"] or record["vsvip_name"]
            ip_values = set(split_multi_value_cell(record["vip_addresses"]))
            for ip_address in ip_values:
                avi.vip_records.append(
                    {
                        "controller": record["controller"],
                        "vsvip_name": record["vsvip_name"],
                        "uuid": record["uuid"],
                        "vip_networks": record["vip_networks"],
                        "IP Address": ip_address,
                    }
                )
                avi.combined_ip_records.append(
                    {
                        "source": record["controller"],
                        "name": record["vsvip_name"],
                        "state": "avi_vip",
                        "ip_address": ip_address,
                        "mac_address": "",
                        "id": entity_id,
                        "network/port_group": record["vip_networks"],
                        "nic_label/device": "avi_vip",
                    }
                )
            avi.vip_counts.update(ip_values)

    virtual_services = avi_reports.get("virtual_services")
    if virtual_services is not None:
        avi.virtual_services_total = int(len(virtual_services.index))

    network_subnets = avi_reports.get("network_subnets")
    if network_subnets is not None:
        avi.network_subnets_total = int(len(network_subnets.index))

    return avi


def collect_detail_sheets(
    workbooks: list[Path],
    avi_network: AviNetworkSummary | None = None,
) -> dict[str, pd.DataFrame]:
    detail_rows: dict[str, list[dict[str, object]]] = {
        "ds_cross_clustered": [],
        "ds_with_no_vms": [],
        "vms_powered_off": [],
        "cluster_configuration_issues": [],
        "host_in_mm": [],
        "host_ht_not_active": [],
        "host_cpu_overprovisioned": [],
        "host_mem_pressure": [],
        "host_power_policy_not_high_perf": [],
        "host_ntp_not_running": [],
        "vms_suspended": [],
        "templates": [],
        "consolidation_needed": [],
        "vms_not_in_connected_state": [],
        "vms_cpu_socket_fix_needed": [],
        "vms_with_cpu_reservation": [],
        "vms_with_memory_reservation": [],
        "vms_with_memory_contention": [],
        "vms_not_using_paravirtual_scsi_adapter": [],
        "vms_network_not_vmxnet3": [],
        "vms_with_duplicate_ip": [],
        "vms_with_duplicate_mac_address": [],
        "duplicate_ips": [],
        "duplicate_mac": [],
        "host_with_duplicate_ip": [],
        "host_with_duplicate_mac": [],
        "avi_se_with_duplicate_ip": [],
        "avi_vip_with_duplicate_ip": [],
        "vms_with_usb": [],
        "vms_with_snapshots": [],
        "vms_hot_add_enabled": [],
        "vms_tools_fix_needed": [],
        "datastores_with_extents_gt_1": [],
        "hw_version_not_latest": [],
        "missing_mandatory_tags": [],
        "vms_with_duplicate_vm_uuid": [],
    }
    uuid_records: list[dict[str, object]] = []
    combined_ip_records: list[dict[str, object]] = []
    combined_mac_records: list[dict[str, object]] = []
    # Rich per-entity-type records, accumulated across all workbooks. Used
    # after the loop to build the host-only/VM-only duplicate detail sheets
    # against the same global, cross-vCenter duplicate IP/MAC sets used for
    # the combined duplicate_ips/duplicate_mac sheets.
    all_host_ip_records: list[dict[str, object]] = []
    all_host_mac_records: list[dict[str, object]] = []
    all_vm_ip_records: list[dict[str, object]] = []
    all_vm_mac_records: list[dict[str, object]] = []

    for workbook_path in workbooks:
        workbook = pd.ExcelFile(workbook_path)
        if "vInfo" not in workbook.sheet_names:
            continue

        vhost = load_sheet(workbook_path, "vHost") if "vHost" in workbook.sheet_names else None
        host_state_map = build_host_state_map(vhost)
        vinfo = load_sheet(workbook_path, "vInfo")
        detail_frame = pd.DataFrame(index=vinfo.index)
        detail_frame["source_vcenter"] = get_normalized_column(vinfo, "VI SDK Server", default="unknown")
        detail_frame["vm_name"] = get_normalized_column(vinfo, "VM", default="")
        detail_frame["vm_id"] = get_normalized_column(vinfo, "VM ID", default="")
        detail_frame["powerstate"] = get_normalized_column(vinfo, "Powerstate", default="unknown")
        detail_frame["template"] = get_normalized_column(vinfo, "Template", default="blank", bool_values=True)
        detail_frame["consolidation_needed"] = get_normalized_column(
            vinfo,
            "Consolidation Needed",
            default="blank",
            bool_values=True,
        )
        detail_frame["connection_state"] = get_normalized_column(vinfo, "Connection state", default="unknown")
        detail_frame["hw_version"] = get_normalized_hw_version_column(vinfo)
        detail_frame["vm_uuid"] = get_normalized_column(vinfo, "VM UUID", default="")
        detail_frame["host"] = get_normalized_column(vinfo, "Host", default="")
        detail_frame["cluster"] = get_normalized_column(vinfo, "Cluster", default="")
        detail_frame["cpus"] = numeric_series(vinfo, "CPUs")
        detail_frame["memory"] = numeric_series(vinfo, "Memory")
        detail_frame["provisioned"] = numeric_series(vinfo, "Provisioned MiB")
        detail_frame["in_use"] = numeric_series(vinfo, "In Use MiB")
        detail_frame["esx_version"] = ""

        cluster_num_hosts_map: dict[str, float] = {}
        if "vCluster" in workbook.sheet_names:
            vcluster = load_sheet(workbook_path, "vCluster")
            if {"Name", "NumHosts"}.issubset(vcluster.columns):
                cluster_num_hosts_map.update(
                    pd.DataFrame(
                        {
                            "cluster_name": get_normalized_column(vcluster, "Name", default=""),
                            "cluster_num_hosts": numeric_series(vcluster, "NumHosts"),
                        }
                    )
                    .loc[lambda df: df["cluster_name"].ne("")]
                    .drop_duplicates(subset=["cluster_name"], keep="first")
                    .set_index("cluster_name")["cluster_num_hosts"]
                    .to_dict()
                )
            if {
                "Name",
                "HA enabled",
                "DRS enabled",
                "DRS default VM behavior",
                "Isolation Response",
            }.issubset(vcluster.columns):
                cluster_rows = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vcluster, "VI SDK Server", default="unknown"),
                        "cluster_name": get_normalized_column(vcluster, "Name", default=""),
                        "ha_enabled": get_normalized_column(
                            vcluster,
                            "HA enabled",
                            default="blank",
                            bool_values=True,
                        ),
                        "drs_enabled": get_normalized_column(
                            vcluster,
                            "DRS enabled",
                            default="blank",
                            bool_values=True,
                        ),
                        "drs_default_vm_behavior": get_normalized_column(
                            vcluster,
                            "DRS default VM behavior",
                            default="unknown",
                        ),
                        "isolation_response": get_normalized_column(
                            vcluster,
                            "Isolation Response",
                            default="unknown",
                        ),
                    }
                ).loc[lambda df: df["cluster_name"].ne("")]
                if not cluster_rows.empty:
                    issue_rows = cluster_rows.loc[
                        lambda df: df["ha_enabled"].eq("false")
                        | df["drs_enabled"].eq("false")
                        | (
                            df["drs_enabled"].eq("true")
                            & df["drs_default_vm_behavior"].ne("fullyAutomated")
                        )
                        | df["isolation_response"].ne("none")
                    ].rename(
                        columns={
                            "cluster_name": "cluster_name",
                            "ha_enabled": "HA enabled",
                            "drs_enabled": "DRS enabled",
                            "drs_default_vm_behavior": "DRS default VM behavior",
                            "isolation_response": "Isolation Response",
                        }
                    )
                    detail_rows["cluster_configuration_issues"].extend(issue_rows.to_dict("records"))
        if not cluster_num_hosts_map and "vHost" in workbook.sheet_names:
            if vhost is not None and {"Host", "Cluster"}.issubset(vhost.columns):
                host_clusters = pd.DataFrame(
                    {
                        "host_name": get_normalized_column(vhost, "Host", default=""),
                        "cluster_name": get_normalized_column(vhost, "Cluster", default=""),
                    }
                ).loc[lambda df: df["host_name"].ne("") & df["cluster_name"].ne("")]
                cluster_num_hosts_map.update(host_clusters.groupby("cluster_name")["host_name"].nunique().to_dict())
        if "vDatastore" in workbook.sheet_names:
            vdatastore = load_sheet(workbook_path, "vDatastore")
            resolved_datastore_frame = build_resolved_datastore_cluster_frame(
                vdatastore,
                build_host_to_cluster_map(vhost) if vhost is not None else {},
            )
            if not resolved_datastore_frame.empty:
                cross_clustered_rows = (
                    resolved_datastore_frame.loc[
                        resolved_datastore_frame["Cluster Count"].gt(1),
                        [
                            "source_vcenter",
                            "datastore_name",
                            "datastore_identifier",
                            "Datastore Locality",
                            "MHA",
                            "Hosts Reported",
                            "Hosts",
                            "Cluster Count",
                            "Cluster Names",
                        ],
                    ]
                    .rename(
                        columns={
                            "Datastore Locality": "datastore_locality",
                            "Hosts Reported": "hosts_reported",
                            "Cluster Count": "cluster_count",
                            "Cluster Names": "cluster_names",
                        }
                    )
                )
                detail_rows["ds_cross_clustered"].extend(cross_clustered_rows.to_dict("records"))
            datastore_rows = build_datastore_inventory_frame(vdatastore)
            if not datastore_rows.empty:
                detail_rows["ds_with_no_vms"].extend(
                    datastore_rows.loc[
                        lambda df: df["is_non_local"] & df["# VMs total"].eq(0),
                        datastore_report_columns(),
                    ].to_dict("records")
                )
                detail_rows["datastores_with_extents_gt_1"].extend(
                    datastore_rows.loc[
                        lambda df: df["# Extents"].gt(1),
                        datastore_report_columns(),
                    ].to_dict("records")
                )
        if "vHost" in workbook.sheet_names:
            if vhost is not None and {"Host", "ESX Version"}.issubset(vhost.columns):
                esx_version_map = (
                    pd.DataFrame(
                        {
                            "host_name": get_normalized_column(vhost, "Host", default=""),
                            "esx_version": get_normalized_column(vhost, "ESX Version", default=""),
                        }
                    )
                    .loc[lambda df: df["host_name"].ne("")]
                    .drop_duplicates(subset=["host_name"], keep="first")
                    .set_index("host_name")["esx_version"]
                    .to_dict()
                )
                detail_frame["esx_version"] = detail_frame["host"].map(esx_version_map).fillna("")
            if vhost is not None and {
                "Host",
                "Cluster",
                "# CPU",
                "# Memory",
                "VM Memory Swapped",
                "VM Memory Ballooned",
                "ESX Version",
                "Current CPU power man. policy",
            }.issubset(vhost.columns):
                host_pressure_rows = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vhost, "VI SDK Server", default="unknown"),
                        "hostname": get_normalized_column(vhost, "Host", default=""),
                        "cluster_name": get_normalized_column(vhost, "Cluster", default=""),
                        "CPU": numeric_series(vhost, "# CPU"),
                        "Memory": numeric_series(vhost, "# Memory"),
                        "swapped": numeric_series(vhost, "VM Memory Swapped"),
                        "ballooned": numeric_series(vhost, "VM Memory Ballooned"),
                        "ESXi Version": get_normalized_column(vhost, "ESX Version", default=""),
                    }
                ).loc[lambda df: df["hostname"].ne("")]
                if not host_pressure_rows.empty:
                    host_pressure_rows = host_pressure_rows.loc[
                        lambda df: df["swapped"].ne(0) | df["ballooned"].ne(0)
                    ]
                    detail_rows["host_mem_pressure"].extend(host_pressure_rows.to_dict("records"))
                power_policy_rows = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vhost, "VI SDK Server", default="unknown"),
                        "hostname": get_normalized_column(vhost, "Host", default=""),
                        "cluster_name": get_normalized_column(vhost, "Cluster", default=""),
                        "CPU": numeric_series(vhost, "# CPU"),
                        "Memory": numeric_series(vhost, "# Memory"),
                        "Current CPU power man. policy": get_normalized_column(
                            vhost,
                            "Current CPU power man. policy",
                            default="unknown",
                        ),
                        "ESXi Version": get_normalized_column(vhost, "ESX Version", default=""),
                    }
                ).loc[lambda df: df["hostname"].ne("")]
                if not power_policy_rows.empty:
                    power_policy_rows = power_policy_rows.loc[
                        lambda df: df["Current CPU power man. policy"].str.lower().ne("high performance")
                    ]
                    detail_rows["host_power_policy_not_high_perf"].extend(power_policy_rows.to_dict("records"))
                ntp_rows = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vhost, "VI SDK Server", default="unknown"),
                        "hostname": get_normalized_column(vhost, "Host", default=""),
                        "cluster_name": get_normalized_column(vhost, "Cluster", default=""),
                        "CPU": numeric_series(vhost, "# CPU"),
                        "Memory": numeric_series(vhost, "# Memory"),
                        "NTPD running": get_normalized_column(
                            vhost,
                            "NTPD running",
                            default="blank",
                            bool_values=True,
                        ),
                        "ESXi Version": get_normalized_column(vhost, "ESX Version", default=""),
                    }
                ).loc[lambda df: df["hostname"].ne("")]
                if not ntp_rows.empty:
                    ntp_rows = ntp_rows.loc[lambda df: df["NTPD running"].ne("true")]
                    detail_rows["host_ntp_not_running"].extend(ntp_rows.to_dict("records"))
            if vhost is not None and {"Host", "Cluster", "# CPU", "# Memory", "in Maintenance Mode", "ESX Version"}.issubset(vhost.columns):
                host_mm_rows = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vhost, "VI SDK Server", default="unknown"),
                        "hostname": get_normalized_column(vhost, "Host", default=""),
                        "cluster_name": get_normalized_column(vhost, "Cluster", default=""),
                        "CPU": numeric_series(vhost, "# CPU"),
                        "Memory": numeric_series(vhost, "# Memory"),
                        "in Maintenance Mode": get_normalized_column(
                            vhost,
                            "in Maintenance Mode",
                            default="blank",
                            bool_values=True,
                        ),
                        "ESXi Version": get_normalized_column(vhost, "ESX Version", default=""),
                    }
                ).loc[lambda df: df["hostname"].ne("")]
                if not host_mm_rows.empty:
                    host_mm_rows = host_mm_rows.loc[lambda df: df["in Maintenance Mode"].eq("true")]
                    detail_rows["host_in_mm"].extend(host_mm_rows.to_dict("records"))
            if vhost is not None and {"Host", "Cluster", "# CPU", "# Memory", "HT Available", "HT Active", "ESX Version"}.issubset(vhost.columns):
                host_ht_rows = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vhost, "VI SDK Server", default="unknown"),
                        "hostname": get_normalized_column(vhost, "Host", default=""),
                        "cluster_name": get_normalized_column(vhost, "Cluster", default=""),
                        "CPU": numeric_series(vhost, "# CPU"),
                        "Memory": numeric_series(vhost, "# Memory"),
                        "HT Available": get_normalized_column(
                            vhost,
                            "HT Available",
                            default="blank",
                            bool_values=True,
                        ),
                        "HT Active": get_normalized_column(
                            vhost,
                            "HT Active",
                            default="blank",
                            bool_values=True,
                        ),
                        "ESXi Version": get_normalized_column(vhost, "ESX Version", default=""),
                    }
                ).loc[lambda df: df["hostname"].ne("")]
                if not host_ht_rows.empty:
                    host_ht_rows = host_ht_rows.loc[
                        lambda df: df["HT Available"].eq("true") & df["HT Active"].ne("true")
                    ]
                    detail_rows["host_ht_not_active"].extend(host_ht_rows.to_dict("records"))
            if vhost is not None and {
                "Host",
                "Cluster",
                "# Cores",
                "# Memory",
                "vCPUs per Core",
                "ESX Version",
            }.issubset(vhost.columns):
                host_cpu_overprovisioned_rows = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vhost, "VI SDK Server", default="unknown"),
                        "hostname": get_normalized_column(vhost, "Host", default=""),
                        "cluster_name": get_normalized_column(vhost, "Cluster", default=""),
                        "Cores": numeric_series(vhost, "# Cores"),
                        "Memory": numeric_series(vhost, "# Memory"),
                        "vCPUs per Core": numeric_series(vhost, "vCPUs per Core"),
                        "ESXi Version": get_normalized_column(vhost, "ESX Version", default=""),
                    }
                ).loc[lambda df: df["hostname"].ne("")]
                if not host_cpu_overprovisioned_rows.empty:
                    host_cpu_overprovisioned_rows = host_cpu_overprovisioned_rows.loc[
                        lambda df: df["vCPUs per Core"].ge(
                            HOST_CPU_OVERPROVISIONED_VCPUS_PER_CORE_THRESHOLD
                        )
                    ]
                    detail_rows["host_cpu_overprovisioned"].extend(
                        host_cpu_overprovisioned_rows.to_dict("records")
                    )

        if "vSC_VMK" in workbook.sheet_names:
            vsc_vmk = load_sheet(workbook_path, "vSC_VMK")
            if {"Host", "IP Address"}.issubset(vsc_vmk.columns):
                host_network_lookup = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vsc_vmk, "VI SDK Server", default="unknown"),
                        "hostname": get_normalized_column(vsc_vmk, "Host", default=""),
                        "Datacenter": get_normalized_column(vsc_vmk, "Datacenter", default=""),
                        "Cluster": get_normalized_column(vsc_vmk, "Cluster", default=""),
                        "Port Group": get_normalized_column(vsc_vmk, "Port Group", default=""),
                        "Device": get_normalized_column(vsc_vmk, "Device", default=""),
                        "Mac Address": get_normalized_column(vsc_vmk, "Mac Address", default=""),
                        "ip_address": get_normalized_column(vsc_vmk, "IP Address", default=""),
                        "Subnet mask": get_normalized_column(vsc_vmk, "Subnet mask", default=""),
                        "Gateway": get_normalized_column(vsc_vmk, "Gateway", default=""),
                        "MTU": numeric_series(vsc_vmk, "MTU"),
                    }
                ).loc[lambda df: df["hostname"].ne("")]
                # Duplicate IPs/MACs are resolved globally (across all
                # workbooks, mixing hosts and VMs) after the main loop below,
                # so this only needs to accumulate the raw rich records here.
                for record in host_network_lookup.to_dict("records"):
                    for ip_address in extract_ipv4_addresses(record["ip_address"]):
                        if is_excluded_duplicate_ip(ip_address):
                            continue
                        all_host_ip_records.append(
                            {
                                "source_vcenter": record["source_vcenter"],
                                "hostname": record["hostname"],
                                "Datacenter": record["Datacenter"],
                                "Cluster": record["Cluster"],
                                "Port Group": record["Port Group"],
                                "Device": record["Device"],
                                "IPv4 Address": ip_address,
                                "Mac Address": record["Mac Address"],
                                "Subnet mask": record["Subnet mask"],
                                "Gateway": record["Gateway"],
                                "MTU": record["MTU"],
                            }
                        )
                        combined_ip_records.append(
                            {
                                "source": record["source_vcenter"],
                                "name": record["hostname"],
                                "state": host_state_map.get(record["hostname"], "unknown"),
                                "ip_address": ip_address,
                                "mac_address": record["Mac Address"],
                                "id": record["hostname"],
                                "network/port_group": record["Port Group"],
                                "nic_label/device": record["Device"],
                            }
                        )
            if {"Host", "Mac Address"}.issubset(vsc_vmk.columns):
                host_mac_frame = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vsc_vmk, "VI SDK Server", default="unknown"),
                        "hostname": get_normalized_column(vsc_vmk, "Host", default=""),
                        "Datacenter": get_normalized_column(vsc_vmk, "Datacenter", default=""),
                        "Cluster": get_normalized_column(vsc_vmk, "Cluster", default=""),
                        "Port Group": get_normalized_column(vsc_vmk, "Port Group", default=""),
                        "Device": get_normalized_column(vsc_vmk, "Device", default=""),
                        "IPv4 Address": get_normalized_column(vsc_vmk, "IP Address", default=""),
                        "Mac Address": get_normalized_column(vsc_vmk, "Mac Address", default=""),
                        "Subnet mask": get_normalized_column(vsc_vmk, "Subnet mask", default=""),
                        "Gateway": get_normalized_column(vsc_vmk, "Gateway", default=""),
                        "MTU": numeric_series(vsc_vmk, "MTU"),
                    }
                ).loc[lambda df: df["hostname"].ne("") & df["Mac Address"].ne("")]
                if not host_mac_frame.empty:
                    all_host_mac_records.extend(host_mac_frame.to_dict("records"))
                    combined_mac_records.extend(
                        {
                            "source": record["source_vcenter"],
                            "name": record["hostname"],
                            "state": host_state_map.get(record["hostname"], "unknown"),
                            "ip_address": record["IPv4 Address"],
                            "mac_address": record["Mac Address"],
                            "id": record["hostname"],
                            "network/port_group": record["Port Group"],
                            "nic_label/device": record["Device"],
                        }
                        for record in host_mac_frame.to_dict("records")
                    )

        # Vectorize the per-column "is this tag missing" check (same rule
        # summarize_workbook() already uses) instead of visiting every cell
        # with a Python-level .at[] lookup; only the final per-row join of
        # missing column names still runs row-wise, over a small boolean
        # frame rather than the raw vInfo sheet.
        missing_tag_mask_frame = pd.DataFrame(
            {
                column: (
                    normalize_strings(vinfo[column], default="").eq("")
                    if column in vinfo.columns
                    else pd.Series(True, index=vinfo.index)
                )
                for column in MANDATORY_TAG_COLUMNS
            },
            index=vinfo.index,
        )
        detail_frame["missing_mandatory_tags"] = missing_tag_mask_frame.apply(
            lambda row: ", ".join(column for column, is_missing in row.items() if is_missing),
            axis=1,
        )

        def add_rows(
            sheet_name: str,
            mask: pd.Series,
            status_column: str,
            output_column_name: str,
        ) -> None:
            if not mask.any():
                return
            subset = detail_frame.loc[mask, ["source_vcenter", "vm_name", "vm_id", status_column]].copy()
            subset = subset.rename(columns={status_column: output_column_name})
            detail_rows[sheet_name].extend(subset.to_dict("records"))

        if detail_frame["powerstate"].eq("poweredOff").any():
            powered_off_rows = detail_frame.loc[
                detail_frame["powerstate"].eq("poweredOff"),
                ["source_vcenter", "vm_name", "vm_id", "powerstate", "cpus", "memory", "provisioned", "in_use"],
            ].rename(
                columns={
                    "powerstate": "Powerstate",
                    "cpus": "CPUs",
                    "memory": "Memory",
                    "provisioned": "Provisioned",
                    "in_use": "In Use",
                }
            )
            detail_rows["vms_powered_off"].extend(powered_off_rows.to_dict("records"))
        add_rows("vms_suspended", detail_frame["powerstate"].eq("suspended"), "powerstate", "Powerstate")
        add_rows("templates", detail_frame["template"].eq("true"), "template", "Template")
        add_rows(
            "consolidation_needed",
            detail_frame["consolidation_needed"].eq("true"),
            "consolidation_needed",
            "Consolidation Needed",
        )
        add_rows(
            "vms_not_in_connected_state",
            detail_frame["connection_state"].ne("connected"),
            "connection_state",
            "Connection state",
        )

        if "vCPU" in workbook.sheet_names and "vHost" in workbook.sheet_names:
            vcpu = load_sheet(workbook_path, "vCPU")
            vhost = load_sheet(workbook_path, "vHost")
            if {"VM ID", "Host", "Sockets"}.issubset(vcpu.columns) and {"Host", "# CPU"}.issubset(vhost.columns):
                host_socket_map = (
                    pd.DataFrame(
                        {
                            "host_name": get_normalized_column(vhost, "Host", default=""),
                            "host_sockets": numeric_series(vhost, "# CPU"),
                        }
                    )
                    .loc[lambda df: df["host_name"].ne("")]
                    .drop_duplicates(subset=["host_name"], keep="first")
                    .set_index("host_name")["host_sockets"]
                    .to_dict()
                )
                vcpu_detail = pd.DataFrame(
                    {
                        "vm_id": get_normalized_column(vcpu, "VM ID", default=""),
                        "host_name": get_normalized_column(vcpu, "Host", default=""),
                        "vm_cpu_sockets": numeric_series(vcpu, "Sockets"),
                    }
                )
                vcpu_detail["host_sockets"] = vcpu_detail["host_name"].map(host_socket_map)

                vm_lookup = detail_frame[["source_vcenter", "vm_name", "vm_id", "cluster"]].copy()
                socket_fix_rows = (
                    vcpu_detail.merge(vm_lookup, on="vm_id", how="left")
                    .loc[
                        lambda df: df["host_name"].ne("")
                        & df["host_sockets"].notna()
                        & df["vm_cpu_sockets"].gt(df["host_sockets"]),
                        ["source_vcenter", "vm_name", "vm_id", "cluster", "host_name", "vm_cpu_sockets", "host_sockets"],
                    ]
                    .rename(columns={"host_name": "Host"})
                )
                socket_fix_rows = socket_fix_rows.rename(columns={"cluster": "Cluster"})
                detail_rows["vms_cpu_socket_fix_needed"].extend(socket_fix_rows.to_dict("records"))

        if "vCPU" in workbook.sheet_names:
            vcpu = load_sheet(workbook_path, "vCPU")
            if {"VM ID", "Reservation"}.issubset(vcpu.columns):
                vcpu_reservation = pd.DataFrame(
                    {
                        "vm_id": get_normalized_column(vcpu, "VM ID", default=""),
                        "reservation": numeric_series(vcpu, "Reservation"),
                    }
                )
                vm_lookup = detail_frame[["source_vcenter", "vm_name", "vm_id"]].copy()
                reservation_rows = (
                    vcpu_reservation.merge(vm_lookup, on="vm_id", how="left")
                    .loc[
                        lambda df: df["reservation"].ne(0),
                        ["source_vcenter", "vm_name", "vm_id", "reservation"],
                    ]
                    .rename(columns={"reservation": "Reservation"})
                )
                detail_rows["vms_with_cpu_reservation"].extend(reservation_rows.to_dict("records"))

            if {"VM ID", "Hot Add"}.issubset(vcpu.columns):
                vcpu_hot_add = pd.DataFrame(
                    {
                        "vm_id": get_normalized_column(vcpu, "VM ID", default=""),
                        "hot_add": get_normalized_column(vcpu, "Hot Add", default="blank", bool_values=True),
                    }
                )
                vm_lookup = detail_frame[["source_vcenter", "vm_name", "vm_id"]].copy()
                hot_add_rows = (
                    vcpu_hot_add.merge(vm_lookup, on="vm_id", how="left")
                    .loc[
                        lambda df: df["hot_add"].eq("true"),
                        ["source_vcenter", "vm_name", "vm_id", "hot_add"],
                    ]
                    .rename(columns={"hot_add": "Hot Add"})
                )
                detail_rows["vms_hot_add_enabled"].extend(hot_add_rows.to_dict("records"))

        if "vMemory" in workbook.sheet_names:
            vmemory = load_sheet(workbook_path, "vMemory")
            if {"VM ID", "Memory Reservation Locked To Max"}.issubset(vmemory.columns):
                vmemory_reservation = pd.DataFrame(
                    {
                        "vm_id": get_normalized_column(vmemory, "VM ID", default=""),
                        "memory_reservation_locked_to_max": get_normalized_column(
                            vmemory,
                            "Memory Reservation Locked To Max",
                            default="blank",
                            bool_values=True,
                        ),
                    }
                )
                vm_lookup = detail_frame[["source_vcenter", "vm_name", "vm_id"]].copy()
                memory_reservation_rows = (
                    vmemory_reservation.merge(vm_lookup, on="vm_id", how="left")
                    .loc[
                        lambda df: df["memory_reservation_locked_to_max"].eq("true"),
                        ["source_vcenter", "vm_name", "vm_id", "memory_reservation_locked_to_max"],
                    ]
                    .rename(
                        columns={
                            "memory_reservation_locked_to_max": "Memory Reservation Locked To Max"
                        }
                    )
                )
                detail_rows["vms_with_memory_reservation"].extend(
                    memory_reservation_rows.to_dict("records")
                )

            if {"VM ID", "Ballooned", "Swapped"}.issubset(vmemory.columns):
                vmemory_contention = pd.DataFrame(
                    {
                        "vm_id": get_normalized_column(vmemory, "VM ID", default=""),
                        "Ballooned": numeric_series(vmemory, "Ballooned"),
                        "Swapped": numeric_series(vmemory, "Swapped"),
                    }
                )
                vmemory_contention["memory_contention"] = (
                    vmemory_contention["Ballooned"] + vmemory_contention["Swapped"]
                )
                vm_lookup = detail_frame[["source_vcenter", "vm_name", "vm_id"]].copy()
                memory_contention_rows = (
                    vmemory_contention.merge(
                        detail_frame[["source_vcenter", "vm_name", "vm_id", "cluster", "host"]],
                        on="vm_id",
                        how="left",
                    )
                    .loc[
                        lambda df: df["memory_contention"].gt(0),
                        ["source_vcenter", "vm_name", "vm_id", "cluster", "host", "Ballooned", "Swapped"],
                    ]
                    .rename(columns={"cluster": "Cluster", "host": "Host"})
                )
                memory_contention_rows["cluster_num_hosts"] = memory_contention_rows["Cluster"].map(
                    cluster_num_hosts_map
                )
                detail_rows["vms_with_memory_contention"].extend(
                    memory_contention_rows.to_dict("records")
                )

        if "vDisk" in workbook.sheet_names:
            vdisk = load_sheet(workbook_path, "vDisk")
            if {"VM ID", "Controller"}.issubset(vdisk.columns):
                vdisk_vm_controller = pd.DataFrame(
                    {
                        "vm_id": get_normalized_column(vdisk, "VM ID", default=""),
                        "controller": get_normalized_column(vdisk, "Controller", default=""),
                    }
                ).loc[lambda df: df["vm_id"].ne("")]
                if not vdisk_vm_controller.empty:
                    controller_by_vm = (
                        vdisk_vm_controller.groupby("vm_id")["controller"]
                        .apply(
                            lambda series: sorted(
                                {value for value in series.tolist() if value}
                            )
                        )
                        .reset_index(name="controllers")
                    )
                    controller_by_vm["uses_paravirtual_scsi"] = controller_by_vm["controllers"].apply(
                        lambda values: "VMware paravirtual SCSI" in values
                    )
                    controller_by_vm["Controller"] = controller_by_vm["controllers"].apply(", ".join)
                    vm_lookup = detail_frame[["source_vcenter", "vm_name", "vm_id"]].copy()
                    non_pvscsi_rows = (
                        controller_by_vm.merge(vm_lookup, on="vm_id", how="left")
                        .loc[
                            lambda df: ~df["uses_paravirtual_scsi"],
                            ["source_vcenter", "vm_name", "vm_id", "Controller"],
                        ]
                    )
                    detail_rows["vms_not_using_paravirtual_scsi_adapter"].extend(
                        non_pvscsi_rows.to_dict("records")
                    )

        if "vNetwork" in workbook.sheet_names:
            vnetwork = load_sheet(workbook_path, "vNetwork")
            if {"VM ID", "Adapter"}.issubset(vnetwork.columns):
                vnetwork_vm_adapter = pd.DataFrame(
                    {
                        "vm_id": get_normalized_column(vnetwork, "VM ID", default=""),
                        "Adapter": get_normalized_column(vnetwork, "Adapter", default=""),
                    }
                ).loc[lambda df: df["vm_id"].ne("")]
                if not vnetwork_vm_adapter.empty:
                    adapter_by_vm = (
                        vnetwork_vm_adapter.groupby("vm_id")["Adapter"]
                        .apply(lambda series: sorted({value for value in series.tolist() if value}))
                        .reset_index(name="adapters")
                    )
                    adapter_by_vm["has_non_vmxnet3"] = adapter_by_vm["adapters"].apply(
                        lambda values: any(value != "Vmxnet3" for value in values)
                    )
                    adapter_by_vm["Adapter"] = adapter_by_vm["adapters"].apply(", ".join)
                    vm_lookup = detail_frame[["source_vcenter", "vm_name", "vm_id"]].copy()
                    non_vmxnet3_rows = (
                        adapter_by_vm.merge(vm_lookup, on="vm_id", how="left")
                        .loc[
                            lambda df: df["has_non_vmxnet3"],
                            ["source_vcenter", "vm_name", "vm_id", "Adapter"],
                        ]
                    )
                    detail_rows["vms_network_not_vmxnet3"].extend(non_vmxnet3_rows.to_dict("records"))
            if {"VM ID", "VM", "IPv4 Address"}.issubset(vnetwork.columns):
                vm_powerstate_map = (
                    detail_frame[["vm_id", "powerstate"]]
                    .drop_duplicates(subset=["vm_id"], keep="first")
                    .set_index("vm_id")["powerstate"]
                    .to_dict()
                )
                base_network_lookup = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vnetwork, "VI SDK Server", default="unknown"),
                        "vm_name": get_normalized_column(vnetwork, "VM", default=""),
                        "vm_id": get_normalized_column(vnetwork, "VM ID", default=""),
                        "NIC label": get_normalized_column(vnetwork, "NIC label", default=""),
                        "Adapter": get_normalized_column(vnetwork, "Adapter", default=""),
                        "Network": get_normalized_column(vnetwork, "Network", default=""),
                        "Switch": get_normalized_column(vnetwork, "Switch", default=""),
                        "Connected": get_normalized_column(vnetwork, "Connected", default=""),
                        "Mac Address": get_normalized_column(vnetwork, "Mac Address", default=""),
                        "ipv4_address": get_normalized_column(vnetwork, "IPv4 Address", default=""),
                    }
                ).loc[lambda df: df["vm_id"].ne("")]
                # Duplicate IPs/MACs are resolved globally (across all
                # workbooks, mixing hosts and VMs) after the main loop below,
                # so this only needs to accumulate the raw rich records here.
                if not base_network_lookup.empty:
                    for record in base_network_lookup.to_dict("records"):
                        for ip_address in extract_ipv4_addresses(record["ipv4_address"]):
                            if is_excluded_duplicate_ip(ip_address):
                                continue
                            all_vm_ip_records.append(
                                {
                                    "source_vcenter": record["source_vcenter"],
                                    "vm_name": record["vm_name"],
                                    "vm_id": record["vm_id"],
                                    "NIC label": record["NIC label"],
                                    "Adapter": record["Adapter"],
                                    "Network": record["Network"],
                                    "Switch": record["Switch"],
                                    "Connected": record["Connected"],
                                    "Powerstate": vm_powerstate_map.get(record["vm_id"], ""),
                                    "IPv4 Address": ip_address,
                                    "Mac Address": record["Mac Address"],
                                }
                            )
                            combined_ip_records.append(
                                {
                                    "source": record["source_vcenter"],
                                    "name": record["vm_name"],
                                    "state": vm_powerstate_map.get(record["vm_id"], "unknown"),
                                    "ip_address": ip_address,
                                    "mac_address": record["Mac Address"],
                                    "id": record["vm_id"],
                                    "network/port_group": record["Network"],
                                    "nic_label/device": record["NIC label"],
                                }
                            )
            if {"VM ID", "VM", "Mac Address"}.issubset(vnetwork.columns):
                vm_powerstate_map = (
                    detail_frame[["vm_id", "powerstate"]]
                    .drop_duplicates(subset=["vm_id"], keep="first")
                    .set_index("vm_id")["powerstate"]
                    .to_dict()
                )
                mac_frame = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vnetwork, "VI SDK Server", default="unknown"),
                        "vm_name": get_normalized_column(vnetwork, "VM", default=""),
                        "vm_id": get_normalized_column(vnetwork, "VM ID", default=""),
                        "NIC label": get_normalized_column(vnetwork, "NIC label", default=""),
                        "Adapter": get_normalized_column(vnetwork, "Adapter", default=""),
                        "Network": get_normalized_column(vnetwork, "Network", default=""),
                        "Switch": get_normalized_column(vnetwork, "Switch", default=""),
                        "Connected": get_normalized_column(vnetwork, "Connected", default=""),
                        "Powerstate": get_normalized_column(vnetwork, "VM ID", default="").map(vm_powerstate_map).fillna(""),
                        "IPv4 Address": get_normalized_column(vnetwork, "IPv4 Address", default=""),
                        "Mac Address": normalize_strings(vnetwork["Mac Address"], default=""),
                    }
                ).loc[lambda df: df["vm_id"].ne("") & df["Mac Address"].ne("")]
                if not mac_frame.empty:
                    all_vm_mac_records.extend(mac_frame.to_dict("records"))
                    combined_mac_records.extend(
                        {
                            "source": record["source_vcenter"],
                            "name": record["vm_name"],
                            "state": record["Powerstate"] or "unknown",
                            "ip_address": record["IPv4 Address"],
                            "mac_address": record["Mac Address"],
                            "id": record["vm_id"],
                            "network/port_group": record["Network"],
                            "nic_label/device": record["NIC label"],
                        }
                        for record in mac_frame.to_dict("records")
                    )

        if "vUSB" in workbook.sheet_names:
            vusb = load_sheet(workbook_path, "vUSB")
            if {"VM ID", "VM"}.issubset(vusb.columns):
                vusb_vm = pd.DataFrame(
                    {
                        "vm_id": get_normalized_column(vusb, "VM ID", default=""),
                        "vm_name": get_normalized_column(vusb, "VM", default=""),
                        "source_vcenter": get_normalized_column(vusb, "VI SDK Server", default="unknown"),
                        "Cluster": get_normalized_column(vusb, "Cluster", default=""),
                        "Host": get_normalized_column(vusb, "Host", default=""),
                        "device_node": get_normalized_column(vusb, "Device Node", default=""),
                        "device_type": get_normalized_column(vusb, "Device Type", default=""),
                    }
                ).loc[lambda df: df["vm_id"].ne("")]
                if not vusb_vm.empty:
                    grouped_usb = (
                        vusb_vm.groupby(["source_vcenter", "vm_name", "vm_id", "Cluster", "Host"], dropna=False)
                        .agg(
                            usb_device_count=("vm_id", "size"),
                            device_nodes=("device_node", lambda s: ", ".join(sorted({v for v in s if v}))),
                            device_types=("device_type", lambda s: ", ".join(sorted({v for v in s if v}))),
                        )
                        .reset_index()
                    )
                    detail_rows["vms_with_usb"].extend(grouped_usb.to_dict("records"))

        if "vSnapshot" in workbook.sheet_names:
            vsnapshot = load_sheet(workbook_path, "vSnapshot")
            if {"VM ID", "VM", "Date / time", "Size MiB (total)"}.issubset(vsnapshot.columns):
                snapshot_rows = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vsnapshot, "VI SDK Server", default="unknown"),
                        "vm_name": get_normalized_column(vsnapshot, "VM", default=""),
                        "vm_id": get_normalized_column(vsnapshot, "VM ID", default=""),
                        "Description": get_normalized_column(vsnapshot, "Description", default=""),
                        "Date / time": datetime_series(vsnapshot, "Date / time"),
                        "Size MiB (total)": numeric_series(vsnapshot, "Size MiB (total)"),
                    }
                ).loc[lambda df: df["vm_id"].ne("")]
                if not snapshot_rows.empty:
                    snapshot_now = datetime.now()
                    snapshot_rows["Snapshot Age (days)"] = (
                        (snapshot_now - snapshot_rows["Date / time"]).dt.total_seconds() / 86400
                    ).round(2)
                    detail_rows["vms_with_snapshots"].extend(snapshot_rows.to_dict("records"))

        if "vTools" in workbook.sheet_names:
            vtools = load_sheet(workbook_path, "vTools")
            if {"VM ID", "VM", "Powerstate", "Tools"}.issubset(vtools.columns):
                tools_rows = pd.DataFrame(
                    {
                        "source_vcenter": get_normalized_column(vtools, "VI SDK Server", default="unknown"),
                        "vm_name": get_normalized_column(vtools, "VM", default=""),
                        "vm_id": get_normalized_column(vtools, "VM ID", default=""),
                        "powerstate": get_normalized_column(vtools, "Powerstate", default="unknown"),
                        "Tools": get_normalized_column(vtools, "Tools", default="unknown"),
                    }
                ).loc[lambda df: df["vm_id"].ne("")]
                if not tools_rows.empty:
                    tools_fix_rows = tools_rows.loc[
                        lambda df: df["powerstate"].eq("poweredOn") & df["Tools"].ne("toolsOk"),
                        ["source_vcenter", "vm_name", "vm_id", "Tools"],
                    ]
                    detail_rows["vms_tools_fix_needed"].extend(tools_fix_rows.to_dict("records"))
        add_rows(
            "hw_version_not_latest",
            detail_frame["hw_version"].ne(LATEST_HW_VERSION),
            "hw_version",
            "HW version",
        )
        if detail_frame["hw_version"].ne(LATEST_HW_VERSION).any():
            hw_not_latest_rows = (
                detail_frame.loc[
                    detail_frame["hw_version"].ne(LATEST_HW_VERSION),
                    ["source_vcenter", "vm_name", "vm_id", "cluster", "host", "esx_version", "hw_version"],
                ]
                .rename(
                    columns={
                        "cluster": "Cluster",
                        "host": "Host",
                        "esx_version": "ESX Version",
                        "hw_version": "HW version",
                    }
                )
            )
            detail_rows["hw_version_not_latest"].extend(hw_not_latest_rows.to_dict("records"))
        add_rows(
            "missing_mandatory_tags",
            detail_frame["missing_mandatory_tags"].ne(""),
            "missing_mandatory_tags",
            "Missing mandatory tags",
        )

        for record in detail_frame.loc[detail_frame["vm_uuid"].ne(""), [
            "source_vcenter",
            "vm_name",
            "vm_id",
            "vm_uuid",
        ]].to_dict("records"):
            uuid_records.append(record)

    # Fold in Avi service engine data IPs / VIP addresses (if --avi-data was
    # supplied) so an IP shared between vSphere and Avi infrastructure is
    # caught by the same duplicate_ips sheet as a vSphere-only collision.
    # Avi entities intentionally don't contribute MAC addresses; see
    # AviNetworkSummary.
    if avi_network is not None:
        combined_ip_records.extend(avi_network.combined_ip_records)

    duplicate_combined_ip_values: set[str] = set()
    if combined_ip_records:
        combined_ip_frame = pd.DataFrame(combined_ip_records).drop_duplicates()
        combined_ip_frame["entity_key"] = (
            combined_ip_frame["source"].astype(str) + "\0" + combined_ip_frame["id"].astype(str)
        )
        duplicate_combined_ip_values = set(
            combined_ip_frame.groupby("ip_address")["entity_key"].nunique().loc[lambda s: s > 1].index
        )
        detail_rows["duplicate_ips"].extend(
            combined_ip_frame.loc[lambda df: df["ip_address"].isin(duplicate_combined_ip_values)]
            .drop_duplicates(
                subset=[
                    "source",
                    "name",
                    "ip_address",
                    "mac_address",
                    "id",
                    "network/port_group",
                    "nic_label/device",
                ]
            )
            .drop(columns=["entity_key"])
            .to_dict("records")
        )

    duplicate_combined_mac_values: set[str] = set()
    if combined_mac_records:
        combined_mac_frame = pd.DataFrame(combined_mac_records).drop_duplicates()
        combined_mac_frame["normalized_mac_address"] = combined_mac_frame["mac_address"].str.lower()
        combined_mac_frame["entity_key"] = (
            combined_mac_frame["source"].astype(str) + "\0" + combined_mac_frame["id"].astype(str)
        )
        duplicate_combined_mac_values = set(
            combined_mac_frame.groupby("normalized_mac_address")["entity_key"]
            .nunique()
            .loc[lambda s: s > 1]
            .index
        )
        detail_rows["duplicate_mac"].extend(
            combined_mac_frame.loc[
                lambda df: df["normalized_mac_address"].isin(duplicate_combined_mac_values)
            ]
            .drop_duplicates(
                subset=[
                    "source",
                    "name",
                    "ip_address",
                    "mac_address",
                    "id",
                    "network/port_group",
                    "nic_label/device",
                ]
            )
            .drop(columns=["normalized_mac_address", "entity_key"])
            .to_dict("records")
        )

    # host_with_duplicate_{ip,mac} and vms_with_duplicate_{ip,mac_address}
    # are filtered from the exact same global, cross-vCenter duplicate
    # IP/MAC sets used above for duplicate_ips/duplicate_mac (so an
    # IP/MAC shared between a host and a VM still counts as a duplicate for
    # both), just restricted to rows for that one entity type.
    if all_host_ip_records and duplicate_combined_ip_values:
        detail_rows["host_with_duplicate_ip"].extend(
            pd.DataFrame(all_host_ip_records)
            .drop_duplicates()
            .loc[lambda df: df["IPv4 Address"].isin(duplicate_combined_ip_values)]
            .to_dict("records")
        )
    if all_vm_ip_records and duplicate_combined_ip_values:
        detail_rows["vms_with_duplicate_ip"].extend(
            pd.DataFrame(all_vm_ip_records)
            .drop_duplicates()
            .loc[lambda df: df["IPv4 Address"].isin(duplicate_combined_ip_values)]
            .to_dict("records")
        )
    if all_host_mac_records and duplicate_combined_mac_values:
        detail_rows["host_with_duplicate_mac"].extend(
            pd.DataFrame(all_host_mac_records)
            .drop_duplicates()
            .loc[lambda df: df["Mac Address"].str.lower().isin(duplicate_combined_mac_values)]
            .to_dict("records")
        )
    if all_vm_mac_records and duplicate_combined_mac_values:
        detail_rows["vms_with_duplicate_mac_address"].extend(
            pd.DataFrame(all_vm_mac_records)
            .drop_duplicates()
            .loc[lambda df: df["Mac Address"].str.lower().isin(duplicate_combined_mac_values)]
            .to_dict("records")
        )
    if avi_network is not None:
        if avi_network.se_ip_records and duplicate_combined_ip_values:
            detail_rows["avi_se_with_duplicate_ip"].extend(
                pd.DataFrame(avi_network.se_ip_records)
                .drop_duplicates()
                .loc[lambda df: df["IP Address"].isin(duplicate_combined_ip_values)]
                .to_dict("records")
            )
        if avi_network.vip_records and duplicate_combined_ip_values:
            detail_rows["avi_vip_with_duplicate_ip"].extend(
                pd.DataFrame(avi_network.vip_records)
                .drop_duplicates()
                .loc[lambda df: df["IP Address"].isin(duplicate_combined_ip_values)]
                .to_dict("records")
            )

    global_uuid_counts = Counter(record["vm_uuid"] for record in uuid_records)
    duplicate_uuid_values = {uuid for uuid, count in global_uuid_counts.items() if count > 1}
    detail_rows["vms_with_duplicate_vm_uuid"].extend(
        {
            "source_vcenter": record["source_vcenter"],
            "vm_name": record["vm_name"],
            "vm_id": record["vm_id"],
            "VM UUID": record["vm_uuid"],
        }
        for record in uuid_records
        if record["vm_uuid"] in duplicate_uuid_values
    )

    detail_frames: dict[str, pd.DataFrame] = {}
    for sheet_name, rows in detail_rows.items():
        frame = pd.DataFrame(rows)
        if frame.empty:
            last_column_name = {
                "ds_cross_clustered": "cluster_names",
                "ds_with_no_vms": "Object ID",
                "vms_powered_off": "Powerstate",
                "cluster_configuration_issues": "Isolation Response",
                "host_in_mm": "in Maintenance Mode",
                "host_ht_not_active": "HT Active",
                "host_cpu_overprovisioned": "vCPUs per Core",
                "host_mem_pressure": "ballooned",
                "host_power_policy_not_high_perf": "Current CPU power man. policy",
                "host_ntp_not_running": "NTPD running",
                "vms_suspended": "Powerstate",
                "templates": "Template",
                "consolidation_needed": "Consolidation Needed",
                "vms_not_in_connected_state": "Connection state",
                "vms_cpu_socket_fix_needed": "vm_cpu_sockets",
                "vms_with_cpu_reservation": "Reservation",
                "vms_with_memory_reservation": "Memory Reservation Locked To Max",
                "vms_with_memory_contention": "memory_contention",
                "vms_not_using_paravirtual_scsi_adapter": "Controller",
                "vms_network_not_vmxnet3": "Adapter",
                "vms_with_duplicate_ip": "IPv4 Address",
                "vms_with_duplicate_mac_address": "Mac Address",
                "duplicate_ips": "ip_address",
                "duplicate_mac": "mac_address",
                "host_with_duplicate_ip": "IPv4 Address",
                "host_with_duplicate_mac": "Mac Address",
                "avi_se_with_duplicate_ip": "IP Address",
                "avi_vip_with_duplicate_ip": "IP Address",
                "vms_with_usb": "usb_device_count",
                "vms_with_snapshots": "Size MiB (total)",
                "vms_hot_add_enabled": "Hot Add",
                "vms_tools_fix_needed": "Tools",
                "datastores_with_extents_gt_1": "# Extents",
                "hw_version_not_latest": "HW version",
                "missing_mandatory_tags": "Missing mandatory tags",
                "vms_with_duplicate_vm_uuid": "VM UUID",
            }[sheet_name]
            if sheet_name == "ds_cross_clustered":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "datastore_name",
                        "datastore_identifier",
                        "datastore_locality",
                        "MHA",
                        "hosts_reported",
                        "Hosts",
                        "cluster_count",
                        "cluster_names",
                    ]
                )
            elif sheet_name == "ds_with_no_vms":
                frame = pd.DataFrame(columns=datastore_report_columns())
            elif sheet_name == "cluster_configuration_issues":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "cluster_name",
                        "HA enabled",
                        "DRS enabled",
                        "DRS default VM behavior",
                        "Isolation Response",
                    ]
                )
            elif sheet_name == "host_mem_pressure":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "hostname",
                        "cluster_name",
                        "CPU",
                        "Memory",
                        "swapped",
                        "ballooned",
                        "ESXi Version",
                    ]
                )
            elif sheet_name == "host_cpu_overprovisioned":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "hostname",
                        "cluster_name",
                        "Cores",
                        "Memory",
                        "vCPUs per Core",
                        "ESXi Version",
                    ]
                )
            elif sheet_name == "host_power_policy_not_high_perf":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "hostname",
                        "cluster_name",
                        "CPU",
                        "Memory",
                        "Current CPU power man. policy",
                        "ESXi Version",
                    ]
                )
            elif sheet_name == "host_ntp_not_running":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "hostname",
                        "cluster_name",
                        "CPU",
                        "Memory",
                        "NTPD running",
                        "ESXi Version",
                    ]
                )
            elif sheet_name == "host_in_mm":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "hostname",
                        "cluster_name",
                        "CPU",
                        "Memory",
                        "in Maintenance Mode",
                        "ESXi Version",
                    ]
                )
            elif sheet_name == "host_ht_not_active":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "hostname",
                        "cluster_name",
                        "CPU",
                        "Memory",
                        "HT Available",
                        "HT Active",
                        "ESXi Version",
                    ]
                )
            elif sheet_name == "vms_cpu_socket_fix_needed":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "vm_name",
                        "vm_id",
                        "Cluster",
                        "Host",
                        "vm_cpu_sockets",
                        "host_sockets",
                    ]
                )
            elif sheet_name == "vms_powered_off":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "vm_name",
                        "vm_id",
                        "Powerstate",
                        "CPUs",
                        "Memory",
                        "Provisioned",
                        "In Use",
                    ]
                )
            elif sheet_name == "vms_with_memory_contention":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "vm_name",
                        "vm_id",
                        "Cluster",
                        "cluster_num_hosts",
                        "Host",
                        "Ballooned",
                        "Swapped",
                    ]
                )
            elif sheet_name == "hw_version_not_latest":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "vm_name",
                        "vm_id",
                        "Cluster",
                        "Host",
                        "ESX Version",
                        "HW version",
                    ]
                )
            elif sheet_name == "vms_with_usb":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "vm_name",
                        "vm_id",
                        "Cluster",
                        "Host",
                        "usb_device_count",
                        "device_nodes",
                        "device_types",
                    ]
                )
            elif sheet_name == "vms_with_snapshots":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "vm_name",
                        "vm_id",
                        "Description",
                        "Date / time",
                        "Snapshot Age (days)",
                        "Size MiB (total)",
                    ]
                )
            elif sheet_name == "datastores_with_extents_gt_1":
                frame = pd.DataFrame(columns=datastore_report_columns())
            elif sheet_name == "vms_with_duplicate_ip":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "vm_name",
                        "vm_id",
                        "NIC label",
                        "Adapter",
                        "Network",
                        "Switch",
                        "Connected",
                        "Powerstate",
                        "IPv4 Address",
                        "Mac Address",
                    ]
                )
            elif sheet_name == "vms_with_duplicate_mac_address":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "vm_name",
                        "vm_id",
                        "NIC label",
                        "Adapter",
                        "Network",
                        "Switch",
                        "Connected",
                        "Powerstate",
                        "IPv4 Address",
                        "Mac Address",
                    ]
                )
            elif sheet_name == "duplicate_ips":
                frame = pd.DataFrame(
                    columns=[
                        "source",
                        "name",
                        "state",
                        "ip_address",
                        "mac_address",
                        "id",
                        "network/port_group",
                        "nic_label/device",
                    ]
                )
            elif sheet_name == "duplicate_mac":
                frame = pd.DataFrame(
                    columns=[
                        "source",
                        "name",
                        "state",
                        "ip_address",
                        "mac_address",
                        "id",
                        "network/port_group",
                        "nic_label/device",
                    ]
                )
            elif sheet_name == "host_with_duplicate_ip":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "hostname",
                        "Datacenter",
                        "Cluster",
                        "Port Group",
                        "Device",
                        "IPv4 Address",
                        "Mac Address",
                        "Subnet mask",
                        "Gateway",
                        "MTU",
                    ]
                )
            elif sheet_name == "host_with_duplicate_mac":
                frame = pd.DataFrame(
                    columns=[
                        "source_vcenter",
                        "hostname",
                        "Datacenter",
                        "Cluster",
                        "Port Group",
                        "Device",
                        "IPv4 Address",
                        "Mac Address",
                        "Subnet mask",
                        "Gateway",
                        "MTU",
                    ]
                )
            elif sheet_name == "avi_se_with_duplicate_ip":
                frame = pd.DataFrame(
                    columns=["controller", "se_name", "uuid", "se_group", "IP Address"]
                )
            elif sheet_name == "avi_vip_with_duplicate_ip":
                frame = pd.DataFrame(
                    columns=["controller", "vsvip_name", "uuid", "vip_networks", "IP Address"]
                )
            else:
                frame = pd.DataFrame(columns=["source_vcenter", "vm_name", "vm_id", last_column_name])
        elif sheet_name == "ds_cross_clustered":
            frame = frame.sort_values(
                by=["cluster_count", "source_vcenter", "datastore_name"],
                ascending=[False, True, True],
                kind="stable",
                na_position="last",
            ).reset_index(drop=True)
        elif sheet_name == "ds_with_no_vms":
            frame = frame.sort_values(
                by=["source_vcenter", "datastore_name", "# Hosts"],
                ascending=[True, True, False],
                kind="stable",
                na_position="last",
            ).reset_index(drop=True)
        elif sheet_name == "datastores_with_extents_gt_1":
            frame = frame.sort_values(
                by=["source_vcenter", "datastore_name", "# Extents"],
                kind="stable",
                na_position="last",
            ).reset_index(drop=True)
        elif sheet_name == "vms_with_duplicate_ip":
            frame = frame.sort_values(
                by=["IPv4 Address", "source_vcenter", "vm_name", "NIC label"],
                kind="stable",
                na_position="last",
            ).reset_index(drop=True)
        elif sheet_name == "vms_with_duplicate_mac_address":
            frame = frame.sort_values(
                by=["Mac Address", "source_vcenter", "vm_name", "NIC label"],
                kind="stable",
                na_position="last",
            ).reset_index(drop=True)
        elif sheet_name == "duplicate_ips":
            frame = frame.sort_values(
                by=["ip_address", "source", "name", "nic_label/device"],
                kind="stable",
                na_position="last",
            ).reset_index(drop=True)
        elif sheet_name == "duplicate_mac":
            frame = frame.sort_values(
                by=["mac_address", "source", "name", "nic_label/device"],
                kind="stable",
                na_position="last",
            ).reset_index(drop=True)
        elif sheet_name == "host_with_duplicate_ip":
            frame = frame.sort_values(
                by=["IPv4 Address", "source_vcenter", "hostname", "Device"],
                kind="stable",
                na_position="last",
            ).reset_index(drop=True)
        elif sheet_name == "host_with_duplicate_mac":
            frame = frame.sort_values(
                by=["Mac Address", "source_vcenter", "hostname", "Device"],
                kind="stable",
                na_position="last",
            ).reset_index(drop=True)
        detail_frames[sheet_name] = frame

    return detail_frames


def summarize_workbook(path: Path) -> WorkbookSummary | None:
    summary = WorkbookSummary(path=path)

    workbook = pd.ExcelFile(path)
    if "vInfo" not in workbook.sheet_names:
        return None

    vinfo = load_sheet(path, "vInfo")
    summary.vm_total = len(vinfo.index)

    if "VI SDK Server" in vinfo.columns:
        vc_series = normalize_strings(vinfo["VI SDK Server"], default="").loc[lambda s: s.ne("")]
        unique_vcenters = sorted(vc_series.unique().tolist())
        if unique_vcenters:
            summary.source_vcenter = ", ".join(unique_vcenters)

    if "vDatastore" in workbook.sheet_names:
        vdatastore = load_sheet(path, "vDatastore")
        host_to_cluster: dict[str, str] = {}
        if "vHost" in workbook.sheet_names:
            host_to_cluster = build_host_to_cluster_map(load_sheet(path, "vHost"))
        resolved_datastore_frame = build_resolved_datastore_cluster_frame(vdatastore, host_to_cluster)
        if not resolved_datastore_frame.empty:
            summary.datastore_total = int(
                resolved_datastore_frame.loc[
                    resolved_datastore_frame["is_non_local"],
                    "datastore_identifier",
                ].nunique()
            )
            summary.vsan_datastore_total = int(
                resolved_datastore_frame.loc[
                    resolved_datastore_frame["is_vsan"],
                    "datastore_identifier",
                ].nunique()
            )
            summary.vsan_datastore_capacity_mib = float(
                resolved_datastore_frame.loc[
                    resolved_datastore_frame["is_vsan"],
                    "Capacity MiB",
                ].sum()
            )
            summary.ds_cross_clustered = int(
                resolved_datastore_frame["Cluster Count"].gt(1).sum()
            )
        datastore_inventory = build_datastore_inventory_frame(vdatastore)
        if not datastore_inventory.empty:
            summary.ds_with_no_vms = int(
                datastore_inventory["is_non_local"].mul(datastore_inventory["# VMs total"].eq(0)).sum()
            )
            summary.vsan_datastore_with_no_vms = int(
                datastore_inventory["is_vsan"].mul(datastore_inventory["# VMs total"].eq(0)).sum()
            )
            summary.datastores_with_extents_gt_1 = int(datastore_inventory["# Extents"].gt(1).sum())

    if "vCluster" in workbook.sheet_names:
        vcluster = load_sheet(path, "vCluster")
        if "Name" in vcluster.columns:
            cluster_names = get_normalized_column(vcluster, "Name", default="").loc[lambda s: s.ne("")]
            summary.cluster_total = int(cluster_names.nunique())
        if {"HA enabled", "DRS enabled", "DRS default VM behavior", "Isolation Response"}.issubset(vcluster.columns):
            ha_enabled = get_normalized_column(vcluster, "HA enabled", default="blank", bool_values=True)
            drs_enabled = get_normalized_column(vcluster, "DRS enabled", default="blank", bool_values=True)
            drs_default_vm_behavior = get_normalized_column(
                vcluster,
                "DRS default VM behavior",
                default="unknown",
            )
            isolation_response = get_normalized_column(
                vcluster,
                "Isolation Response",
                default="unknown",
            )
            summary.clusters_ha_disabled = int(ha_enabled.eq("false").sum())
            summary.clusters_drs_disabled = int(drs_enabled.eq("false").sum())
            summary.clusters_drs_not_fullyautomatic = int(
                drs_enabled.eq("true").mul(drs_default_vm_behavior.ne("fullyAutomated")).sum()
            )
            summary.host_isolation_response_not_none = int(isolation_response.ne("none").sum())

    if "vHost" in workbook.sheet_names:
        vhost = load_sheet(path, "vHost")
        if "Host" in vhost.columns:
            hosts = get_normalized_column(vhost, "Host", default="").loc[lambda s: s.ne("")]
            summary.host_total = int(hosts.nunique())
        if "in Maintenance Mode" in vhost.columns:
            summary.host_in_mm = int(
                get_normalized_column(
                    vhost,
                    "in Maintenance Mode",
                    default="blank",
                    bool_values=True,
                ).eq("true").sum()
            )
        if "CPU Model" in vhost.columns:
            cpu_models = get_normalized_column(vhost, "CPU Model", default="").loc[lambda s: s.ne("")]
            summary.host_cpu_model_counts.update(cpu_models.tolist())
        if "ESX Version" in vhost.columns:
            esxi_versions = get_normalized_column(vhost, "ESX Version", default="").loc[lambda s: s.ne("")]
            summary.host_esxi_version_counts.update(esxi_versions.tolist())
            summary.host_esxi_versions = int(len(summary.host_esxi_version_counts))
        if "Vendor" in vhost.columns:
            vendors = get_normalized_column(vhost, "Vendor", default="").loc[lambda s: s.ne("")]
            summary.host_vendor_counts.update(vendors.tolist())
            summary.host_vendors = int(len(summary.host_vendor_counts))
        if "Model" in vhost.columns:
            models = get_normalized_column(vhost, "Model", default="").loc[lambda s: s.ne("")]
            summary.host_model_counts.update(models.tolist())
            summary.host_models = int(len(summary.host_model_counts))
        if "# Cores" in vhost.columns:
            summary.host_cpu_cores = int(numeric_series(vhost, "# Cores").sum())
        if "# Memory" in vhost.columns:
            summary.host_memory_mib = float(numeric_series(vhost, "# Memory").sum())
        if {"HT Available", "HT Active"}.issubset(vhost.columns):
            ht_available = get_normalized_column(vhost, "HT Available", default="blank", bool_values=True)
            ht_active = get_normalized_column(vhost, "HT Active", default="blank", bool_values=True)
            summary.host_ht_not_active = int(ht_available.eq("true").mul(ht_active.ne("true")).sum())
        if "vCPUs per Core" in vhost.columns:
            summary.host_cpu_overprovisioned = int(
                numeric_series(vhost, "vCPUs per Core")
                .ge(HOST_CPU_OVERPROVISIONED_VCPUS_PER_CORE_THRESHOLD)
                .sum()
            )
        if "Current CPU power man. policy" in vhost.columns:
            current_power_policy = get_normalized_column(vhost, "Current CPU power man. policy", default="unknown")
            summary.host_power_policy_not_high_perf = int(
                current_power_policy.str.lower().ne("high performance").sum()
            )
        if "NTPD running" in vhost.columns:
            summary.host_ntp_not_running = int(
                get_normalized_column(vhost, "NTPD running", default="blank", bool_values=True).ne("true").sum()
            )
        if {"VM Memory Swapped", "VM Memory Ballooned"}.issubset(vhost.columns):
            host_vm_memory_swapped = numeric_series(vhost, "VM Memory Swapped")
            host_vm_memory_ballooned = numeric_series(vhost, "VM Memory Ballooned")
            summary.host_mem_pressure = int(
                host_vm_memory_swapped.ne(0).mul(1).add(host_vm_memory_ballooned.ne(0).mul(1)).gt(0).sum()
            )

    if "Powerstate" in vinfo.columns:
        power_states = normalize_strings(vinfo["Powerstate"])
        summary.power_states.update(power_states)
        summary.powered_on_vms = int(power_states.eq("poweredOn").sum())
        summary.powered_off_vms = int(power_states.eq("poweredOff").sum())
        summary.suspended_vms = int(power_states.eq("suspended").sum())

    cpu_values = numeric_series(vinfo, "CPUs")
    summary.vm_vcpu_total = int(cpu_values.sum())
    summary.cpu_le_8 = int(cpu_values.le(8).sum())
    summary.cpu_gt_8_le_16 = int(cpu_values.gt(8).mul(cpu_values.le(16)).sum())
    summary.cpu_gt_16_le_32 = int(cpu_values.gt(16).mul(cpu_values.le(32)).sum())
    summary.cpu_gt_32_le_64 = int(cpu_values.gt(32).mul(cpu_values.le(64)).sum())
    summary.cpu_gt_64 = int(cpu_values.gt(64).sum())

    memory_values = numeric_series(vinfo, "Memory")
    summary.vm_memory_mib = float(memory_values.sum())
    summary.memory_le_16_gib = int(memory_values.le(16 * MIB_PER_GIB).sum())
    summary.memory_gt_16_le_32_gib = int(
        memory_values.gt(16 * MIB_PER_GIB).mul(memory_values.le(32 * MIB_PER_GIB)).sum()
    )
    summary.memory_gt_32_le_64_gib = int(
        memory_values.gt(32 * MIB_PER_GIB).mul(memory_values.le(64 * MIB_PER_GIB)).sum()
    )
    summary.memory_gt_64_le_128_gib = int(
        memory_values.gt(64 * MIB_PER_GIB).mul(memory_values.le(128 * MIB_PER_GIB)).sum()
    )
    summary.memory_gt_128_gib = int(memory_values.gt(128 * MIB_PER_GIB).sum())
    summary.provisioned_mib = float(numeric_series(vinfo, "Provisioned MiB").sum())
    summary.used_mib = float(numeric_series(vinfo, "In Use MiB").sum())
    poweroff_storage_mask = get_normalized_column(vinfo, "Powerstate", default="unknown").eq("poweredOff")
    summary.poweroff_provisioned_mib = float(numeric_series(vinfo, "Provisioned MiB").loc[poweroff_storage_mask].sum())
    summary.poweroff_used_mib = float(numeric_series(vinfo, "In Use MiB").loc[poweroff_storage_mask].sum())

    if "Template" in vinfo.columns:
        summary.template_vms = int(normalize_bool_strings(vinfo["Template"]).eq("true").sum())

    if "SRM Placeholder" in vinfo.columns:
        summary.srm_placeholder_vms = int(
            normalize_bool_strings(vinfo["SRM Placeholder"]).eq("true").sum()
        )

    if "Consolidation Needed" in vinfo.columns:
        summary.consolidation_needed_vms = int(
            normalize_bool_strings(vinfo["Consolidation Needed"]).eq("true").sum()
        )

    if "Connection state" in vinfo.columns:
        connection_states = normalize_strings(vinfo["Connection state"])
        summary.connection_states.update(connection_states)
        summary.vms_not_connected = int(connection_states.ne("connected").sum())

    if "EFI Secure boot" in vinfo.columns:
        summary.efi_secure_boot.update(normalize_bool_strings(vinfo["EFI Secure boot"]))

    if "HW version" in vinfo.columns:
        hw_versions = get_normalized_hw_version_column(vinfo)
        summary.hw_versions.update(hw_versions)
        summary.hw_version_not_latest_vms = int(hw_versions.ne(LATEST_HW_VERSION).sum())

    missing_tag_mask = pd.Series(False, index=vinfo.index)
    for column in MANDATORY_TAG_COLUMNS:
        if column not in vinfo.columns:
            summary.missing_tag_counts[column] += summary.vm_total
            missing_tag_mask |= True
            continue

        missing_for_column = normalize_strings(vinfo[column], default="").eq("")
        summary.missing_tag_counts[column] += int(missing_for_column.sum())
        missing_tag_mask |= missing_for_column

    summary.vms_missing_mandatory_tags = int(missing_tag_mask.sum())

    if "VM UUID" in vinfo.columns:
        vm_uuid = normalize_strings(vinfo["VM UUID"], default="")
        populated_uuid = vm_uuid[vm_uuid.ne("")]
        summary.vm_uuid_counts.update(populated_uuid.tolist())

    if "vCPU" in workbook.sheet_names and "vHost" in workbook.sheet_names:
        vcpu = load_sheet(path, "vCPU")
        vhost = load_sheet(path, "vHost")
        if {"Host", "Sockets"}.issubset(vcpu.columns) and {"Host", "# CPU"}.issubset(vhost.columns):
            host_socket_map = (
                vhost.assign(
                    host_name=normalize_strings(vhost["Host"], default=""),
                    host_cpu=numeric_series(vhost, "# CPU"),
                )
                .loc[lambda df: df["host_name"].ne(""), ["host_name", "host_cpu"]]
                .drop_duplicates(subset=["host_name"], keep="first")
                .set_index("host_name")["host_cpu"]
                .to_dict()
            )
            vcpu_hosts = normalize_strings(vcpu["Host"], default="")
            vcpu_sockets = numeric_series(vcpu, "Sockets")
            summary.vms_cpu_socket_fix_needed = int(
                sum(
                    1
                    for host_name, socket_count in zip(vcpu_hosts, vcpu_sockets)
                    if host_name in host_socket_map and socket_count > host_socket_map[host_name]
                )
            )
            summary.vms_with_cpu_reservation = int(numeric_series(vcpu, "Reservation").ne(0).sum())
            summary.vms_hot_add_enabled = int(
                get_normalized_column(vcpu, "Hot Add", default="blank", bool_values=True).eq("true").sum()
            )
    elif "vCPU" in workbook.sheet_names:
        vcpu = load_sheet(path, "vCPU")
        summary.vms_with_cpu_reservation = int(numeric_series(vcpu, "Reservation").ne(0).sum())
        summary.vms_hot_add_enabled = int(
            get_normalized_column(vcpu, "Hot Add", default="blank", bool_values=True).eq("true").sum()
        )

    if "vMemory" in workbook.sheet_names:
        vmemory = load_sheet(path, "vMemory")
        summary.vms_with_memory_reservation = int(
            get_normalized_column(
                vmemory,
                "Memory Reservation Locked To Max",
                default="blank",
                bool_values=True,
            ).eq("true").sum()
        )
        ballooned = numeric_series(vmemory, "Ballooned")
        swapped = numeric_series(vmemory, "Swapped")
        summary.vms_mem_ballooned = int(ballooned.gt(0).sum())
        summary.vms_mem_swapped = int(swapped.gt(0).sum())
        summary.vms_with_memory_contention = int((ballooned.add(swapped)).gt(0).sum())

    if "vTools" in workbook.sheet_names:
        vtools = load_sheet(path, "vTools")
        if {"Powerstate", "Tools"}.issubset(vtools.columns):
            tools_powerstate = get_normalized_column(vtools, "Powerstate", default="unknown")
            tools_status = get_normalized_column(vtools, "Tools", default="unknown")
            summary.vms_tools_fix_needed = int(
                tools_powerstate.eq("poweredOn").mul(tools_status.ne("toolsOk")).sum()
            )

    if "vDisk" in workbook.sheet_names:
        vdisk = load_sheet(path, "vDisk")
        if {"VM ID", "Controller"}.issubset(vdisk.columns):
            vdisk_vm_controller = pd.DataFrame(
                {
                    "vm_id": get_normalized_column(vdisk, "VM ID", default=""),
                    "controller": get_normalized_column(vdisk, "Controller", default=""),
                }
            ).loc[lambda df: df["vm_id"].ne("") & df["controller"].ne("")]
            if not vdisk_vm_controller.empty:
                unique_vm_controller = vdisk_vm_controller.drop_duplicates(subset=["vm_id", "controller"])
                summary.controller_type_counts.update(unique_vm_controller["controller"].tolist())
                controller_by_vm = vdisk_vm_controller.groupby("vm_id")["controller"].apply(
                    lambda series: {value for value in series.tolist() if value}
                )
                summary.vms_not_using_paravirtual_scsi_adapter = int(
                    sum(
                        1
                        for controllers_for_vm in controller_by_vm
                        if "VMware paravirtual SCSI" not in controllers_for_vm
                    )
                )

    if "vNetwork" in workbook.sheet_names:
        vnetwork = load_sheet(path, "vNetwork")
        if {"VM ID", "Adapter"}.issubset(vnetwork.columns):
            vnetwork_vm_adapter = pd.DataFrame(
                {
                    "vm_id": get_normalized_column(vnetwork, "VM ID", default=""),
                    "adapter": get_normalized_column(vnetwork, "Adapter", default=""),
                }
            ).loc[lambda df: df["vm_id"].ne("") & df["adapter"].ne("")]
            if not vnetwork_vm_adapter.empty:
                unique_vm_adapter = vnetwork_vm_adapter.drop_duplicates(subset=["vm_id", "adapter"])
                summary.network_adapter_type_counts.update(unique_vm_adapter["adapter"].tolist())
                adapter_by_vm = vnetwork_vm_adapter.groupby("vm_id")["adapter"].apply(
                    lambda series: {value for value in series.tolist() if value}
                )
                summary.vms_network_not_vmxnet3 = int(
                    sum(
                        1
                        for adapters_for_vm in adapter_by_vm
                        if any(value != "Vmxnet3" for value in adapters_for_vm)
                    )
                )
        # summary.vms_with_duplicate_ip/vms_with_duplicate_mac_address are
        # resolved later, across all workbooks together, by
        # apply_global_duplicate_network_counts() -- an IP/MAC reused by two
        # VMs (or a VM and a host) in two different vCenter exports is still
        # the same conflict on the wire, so duplicate detection is
        # intentionally not scoped by source vCenter. This just needs to
        # record, per distinct VM, which values it uses.
        if {"VM ID", "IPv4 Address"}.issubset(vnetwork.columns):
            vm_ip_source = pd.DataFrame(
                {
                    "vm_id": get_normalized_column(vnetwork, "VM ID", default=""),
                    "ipv4_address": get_normalized_column(vnetwork, "IPv4 Address", default=""),
                }
            ).loc[lambda df: df["vm_id"].ne("")]
            vm_ip_pairs: set[tuple[str, str]] = set()
            for record in vm_ip_source.to_dict("records"):
                for ip_address in extract_ipv4_addresses(record["ipv4_address"]):
                    if is_excluded_duplicate_ip(ip_address):
                        continue
                    vm_ip_pairs.add((record["vm_id"], ip_address))
            summary.vm_ip_counts.update(ip_address for _, ip_address in vm_ip_pairs)
        if {"VM ID", "Mac Address"}.issubset(vnetwork.columns):
            mac_frame = pd.DataFrame(
                {
                    "vm_id": get_normalized_column(vnetwork, "VM ID", default=""),
                    "mac_address": normalize_strings(vnetwork["Mac Address"], default="").str.lower(),
                }
            ).loc[lambda df: df["vm_id"].ne("") & df["mac_address"].ne("")]
            if not mac_frame.empty:
                unique_vm_mac = mac_frame.drop_duplicates(subset=["vm_id", "mac_address"])
                summary.vm_mac_counts.update(unique_vm_mac["mac_address"].tolist())

    if "vSC_VMK" in workbook.sheet_names:
        vsc_vmk = load_sheet(path, "vSC_VMK")
        if {"Host", "IP Address"}.issubset(vsc_vmk.columns):
            host_ip_source = pd.DataFrame(
                {
                    "hostname": get_normalized_column(vsc_vmk, "Host", default=""),
                    "ip_address": get_normalized_column(vsc_vmk, "IP Address", default=""),
                }
            ).loc[lambda df: df["hostname"].ne("")]
            host_ip_pairs: set[tuple[str, str]] = set()
            for record in host_ip_source.to_dict("records"):
                for ip_address in extract_ipv4_addresses(record["ip_address"]):
                    if is_excluded_duplicate_ip(ip_address):
                        continue
                    host_ip_pairs.add((record["hostname"], ip_address))
            summary.host_ip_counts.update(ip_address for _, ip_address in host_ip_pairs)
        if {"Host", "Mac Address"}.issubset(vsc_vmk.columns):
            host_mac_frame = pd.DataFrame(
                {
                    "hostname": get_normalized_column(vsc_vmk, "Host", default=""),
                    "mac_address": get_normalized_column(vsc_vmk, "Mac Address", default="").str.lower(),
                }
            ).loc[lambda df: df["hostname"].ne("") & df["mac_address"].ne("")]
            if not host_mac_frame.empty:
                unique_host_mac = host_mac_frame.drop_duplicates(subset=["hostname", "mac_address"])
                summary.host_mac_counts.update(unique_host_mac["mac_address"].tolist())

    if "vUSB" in workbook.sheet_names:
        vusb = load_sheet(path, "vUSB")
        if "VM ID" in vusb.columns:
            summary.vms_with_usb = int(get_normalized_column(vusb, "VM ID", default="").loc[lambda s: s.ne("")].nunique())

    if "vSnapshot" in workbook.sheet_names:
        vsnapshot = load_sheet(path, "vSnapshot")
        if {"VM ID", "Date / time"}.issubset(vsnapshot.columns):
            snapshot_vm = pd.DataFrame(
                {
                    "vm_id": get_normalized_column(vsnapshot, "VM ID", default=""),
                    "snapshot_time": datetime_series(vsnapshot, "Date / time"),
                }
            ).loc[lambda df: df["vm_id"].ne("")]
            if not snapshot_vm.empty:
                oldest_snapshot_by_vm = snapshot_vm.groupby("vm_id")["snapshot_time"].min().dropna()
                if not oldest_snapshot_by_vm.empty:
                    now = datetime.now()
                    snapshot_age_hours = (now - oldest_snapshot_by_vm).dt.total_seconds() / 3600
                    summary.vms_with_snapshot = int(len(oldest_snapshot_by_vm))
                    summary.vms_with_snapshot_age_le_48h = int(snapshot_age_hours.le(48).sum())
                    summary.vms_with_snapshot_age_gt_48h_le_7d = int(
                        snapshot_age_hours.gt(48).mul(snapshot_age_hours.le(24 * 7)).sum()
                    )
                    summary.vms_with_snapshot_age_gt_7d = int(snapshot_age_hours.gt(24 * 7).sum())

    return summary


def format_counter(counter: Counter[str]) -> str:
    if not counter:
        return "none"
    return ", ".join(f"{key}={value}" for key, value in sorted(counter.items()))


def mib_to_tb(value: float) -> float:
    return round(float(value) / MIB_PER_TB, 2)


def format_fields(fields: list[tuple[str, object]]) -> list[str]:
    if not fields:
        return []
    width = max(len(name) for name, _ in fields)
    return [f"  {name:<{width}} : {value}" for name, value in fields]


def print_fields(fields: list[tuple[str, object]]) -> None:
    for line in format_fields(fields):
        print(line)


# Single source of truth for the per-workbook metrics shared by the CSV/Excel
# summary row, the per-workbook console summary, and the combined console
# summary. Previously these ~70 fields were hand-duplicated across
# summary_to_row/summary_fields/aggregate_fields, so adding a metric meant
# remembering to update three places -- easy to do inconsistently. Every
# consumer now derives from this list instead.
SUMMARY_FIELD_DEFINITIONS: list[tuple[str, Callable[[WorkbookSummary], object]]] = [
    ("source_vcenter", lambda s: s.source_vcenter),
    ("datastores_total", lambda s: s.datastore_total),
    ("vsan_datastores_total", lambda s: s.vsan_datastore_total),
    ("vsan_datastores_with_no_vms", lambda s: s.vsan_datastore_with_no_vms),
    ("vsan_datastores_capacity_tb", lambda s: mib_to_tb(s.vsan_datastore_capacity_mib)),
    ("ds_cross_clustered", lambda s: s.ds_cross_clustered),
    ("ds_with_no_vms", lambda s: s.ds_with_no_vms),
    ("clusters_total", lambda s: s.cluster_total),
    ("clusters_ha_disabled", lambda s: s.clusters_ha_disabled),
    ("clusters_drs_disabled", lambda s: s.clusters_drs_disabled),
    ("clusters_drs_not_fullyautomatic", lambda s: s.clusters_drs_not_fullyautomatic),
    ("host_isolation_response_not_none", lambda s: s.host_isolation_response_not_none),
    ("hosts_total", lambda s: s.host_total),
    ("host_in_mm", lambda s: s.host_in_mm),
    ("host_ht_not_active", lambda s: s.host_ht_not_active),
    ("host_cpu_overprovisioned", lambda s: s.host_cpu_overprovisioned),
    ("host_mem_pressure", lambda s: s.host_mem_pressure),
    ("host_power_policy_not_high_perf", lambda s: s.host_power_policy_not_high_perf),
    ("host_esxi_versions", lambda s: s.host_esxi_versions),
    ("host_ntp_not_running", lambda s: s.host_ntp_not_running),
    ("host_vendors", lambda s: s.host_vendors),
    ("host_models", lambda s: s.host_models),
    ("host_cpu_cores", lambda s: s.host_cpu_cores),
    ("host_memory", lambda s: mib_to_tb(s.host_memory_mib)),
    ("host_with_duplicate_ip", lambda s: s.host_with_duplicate_ip),
    ("host_with_duplicate_mac", lambda s: s.host_with_duplicate_mac),
    ("duplicate_ips", lambda s: s.duplicate_ips),
    ("duplicate_mac", lambda s: s.duplicate_mac),
    ("host_cpu_models", lambda s: len(s.host_cpu_model_counts)),
    ("vms_total", lambda s: s.vm_total),
    ("vms_vcpu_total", lambda s: s.vm_vcpu_total),
    ("vms_memory_total", lambda s: mib_to_tb(s.vm_memory_mib)),
    ("vms_powered_on", lambda s: s.powered_on_vms),
    ("vms_powered_off", lambda s: s.powered_off_vms),
    ("vms_suspended", lambda s: s.suspended_vms),
    ("vms_templates", lambda s: s.template_vms),
    ("vms_srm_placeholder", lambda s: s.srm_placeholder_vms),
    ("vms_consolidation_needed", lambda s: s.consolidation_needed_vms),
    ("vms_in_not_connected_state", lambda s: s.vms_not_connected),
    ("vms_cpu_socket_fix_needed", lambda s: s.vms_cpu_socket_fix_needed),
    ("vms_with_cpu_reservation", lambda s: s.vms_with_cpu_reservation),
    ("vms_with_memory_reservation", lambda s: s.vms_with_memory_reservation),
    ("vms_mem_ballooned", lambda s: s.vms_mem_ballooned),
    ("vms_mem_swapped", lambda s: s.vms_mem_swapped),
    ("vms_with_memory_contention", lambda s: s.vms_with_memory_contention),
    ("vms_not_using_paravirtual_scsi_adapter", lambda s: s.vms_not_using_paravirtual_scsi_adapter),
    ("vms_network_not_vmxnet3", lambda s: s.vms_network_not_vmxnet3),
    ("vms_with_duplicate_ip", lambda s: s.vms_with_duplicate_ip),
    ("vms_with_duplicate_mac_address", lambda s: s.vms_with_duplicate_mac_address),
    ("vms_with_usb", lambda s: s.vms_with_usb),
    ("vms_with_snapshot", lambda s: s.vms_with_snapshot),
    ("vms_with_snapshot_age_le_48h", lambda s: s.vms_with_snapshot_age_le_48h),
    ("vms_with_snapshot_age_gt_48h_le_7d", lambda s: s.vms_with_snapshot_age_gt_48h_le_7d),
    ("vms_with_snapshot_age_gt_7d", lambda s: s.vms_with_snapshot_age_gt_7d),
    ("vms_hot_add_enabled", lambda s: s.vms_hot_add_enabled),
    ("vms_tools_fix_needed", lambda s: s.vms_tools_fix_needed),
    ("datastores_with_extents_gt_1", lambda s: s.datastores_with_extents_gt_1),
    ("vms_hw_version_not_latest", lambda s: s.hw_version_not_latest_vms),
    ("vms_missing_mandatory_tags", lambda s: s.vms_missing_mandatory_tags),
    ("vms_with_duplicate_vm_uuid", lambda s: s.duplicate_uuid_vms),
    ("vms_cpu_le_8", lambda s: s.cpu_le_8),
    ("vms_cpu_gt_8_le_16", lambda s: s.cpu_gt_8_le_16),
    ("vms_cpu_gt_16_le_32", lambda s: s.cpu_gt_16_le_32),
    ("vms_cpu_gt_32_le_64", lambda s: s.cpu_gt_32_le_64),
    ("vms_cpu_gt_64", lambda s: s.cpu_gt_64),
    ("vms_memory_le_16_gib", lambda s: s.memory_le_16_gib),
    ("vms_memory_gt_16_le_32_gib", lambda s: s.memory_gt_16_le_32_gib),
    ("vms_memory_gt_32_le_64_gib", lambda s: s.memory_gt_32_le_64_gib),
    ("vms_memory_gt_64_le_128_gib", lambda s: s.memory_gt_64_le_128_gib),
    ("vms_memory_gt_128_gib", lambda s: s.memory_gt_128_gib),
    ("vms_storage_provisioned", lambda s: mib_to_tb(s.provisioned_mib)),
    ("vms_storage_used", lambda s: mib_to_tb(s.used_mib)),
    ("vms_storage_provisioned_poweroff", lambda s: mib_to_tb(s.poweroff_provisioned_mib)),
    ("vms_storage_used_poweroff", lambda s: mib_to_tb(s.poweroff_used_mib)),
]


def base_summary_fields(summary: WorkbookSummary) -> list[tuple[str, object]]:
    return [(key, extractor(summary)) for key, extractor in SUMMARY_FIELD_DEFINITIONS]


def controller_adapter_fields(
    summary: WorkbookSummary,
    controller_keys: list[str] | None = None,
    adapter_keys: list[str] | None = None,
) -> list[tuple[str, object]]:
    fields: list[tuple[str, object]] = list(
        controller_count_columns(
            summary.controller_type_counts,
            controller_keys if controller_keys is not None else sorted(summary.controller_type_counts),
        ).items()
    )
    fields.extend(
        adapter_count_columns(
            summary.network_adapter_type_counts,
            adapter_keys if adapter_keys is not None else sorted(summary.network_adapter_type_counts),
        ).items()
    )
    return fields


def counter_breakdown_fields(summary: WorkbookSummary) -> list[tuple[str, object]]:
    return [
        ("vms_connection_state", format_counter(summary.connection_states)),
        ("vms_efi_secure_boot", format_counter(summary.efi_secure_boot)),
        ("vms_hw_version", format_counter(summary.hw_versions)),
        ("vms_missing_tags_breakdown", format_counter(summary.missing_tag_counts)),
    ]


def summary_to_row(
    summary: WorkbookSummary,
    controller_keys: list[str],
    adapter_keys: list[str],
) -> dict[str, object]:
    row = dict(base_summary_fields(summary))
    row.update(controller_adapter_fields(summary, controller_keys, adapter_keys))
    row.update(counter_breakdown_fields(summary))
    row["source_path"] = str(summary.path)
    return row


def _reconcile_duplicate_network_counts(
    summary: WorkbookSummary,
    duplicate_ip_values: set[str],
    duplicate_mac_values: set[str],
) -> None:
    """Set summary.{host,vms}_with_duplicate_{ip,mac} and duplicate_{ips,mac}
    from summary's own host_ip_counts/vm_ip_counts/host_mac_counts/
    vm_mac_counts, given the sets of values already known to be duplicated
    globally (across all workbooks, mixing hosts and VMs). duplicate_ips and
    duplicate_mac are simply the host + vm figures added together, since
    every entity counted there is either a host or a VM.
    """
    summary.host_with_duplicate_ip = sum(
        count for value, count in summary.host_ip_counts.items() if value in duplicate_ip_values
    )
    summary.vms_with_duplicate_ip = sum(
        count for value, count in summary.vm_ip_counts.items() if value in duplicate_ip_values
    )
    summary.duplicate_ips = summary.host_with_duplicate_ip + summary.vms_with_duplicate_ip
    summary.host_with_duplicate_mac = sum(
        count for value, count in summary.host_mac_counts.items() if value in duplicate_mac_values
    )
    summary.vms_with_duplicate_mac_address = sum(
        count for value, count in summary.vm_mac_counts.items() if value in duplicate_mac_values
    )
    summary.duplicate_mac = summary.host_with_duplicate_mac + summary.vms_with_duplicate_mac_address


def apply_global_duplicate_network_counts(
    summaries: list[WorkbookSummary],
    extra_ip_counts: Counter[str] | None = None,
    extra_mac_counts: Counter[str] | None = None,
) -> tuple[set[str], set[str]]:
    """Resolve host_with_duplicate_ip/mac and vms_with_duplicate_ip/mac
    (and duplicate_ips/duplicate_mac) across all workbooks together, and
    return the resulting (duplicate_ip_values, duplicate_mac_values) sets so
    callers can reuse them for the combined summary and for Avi's own
    duplicate counts.

    An IP or MAC address reused by two entities is the same real-world
    conflict whether those entities came from the same RVTools export, from
    two different vCenters, or from Avi infrastructure (extra_ip_counts/
    extra_mac_counts), so this intentionally does not scope duplicate
    detection by source vCenter or by vSphere-vs-Avi origin.
    """
    global_host_ip: Counter[str] = Counter()
    global_vm_ip: Counter[str] = Counter()
    global_host_mac: Counter[str] = Counter()
    global_vm_mac: Counter[str] = Counter()
    for summary in summaries:
        global_host_ip.update(summary.host_ip_counts)
        global_vm_ip.update(summary.vm_ip_counts)
        global_host_mac.update(summary.host_mac_counts)
        global_vm_mac.update(summary.vm_mac_counts)

    combined_ip_counts = global_host_ip + global_vm_ip
    if extra_ip_counts:
        combined_ip_counts += extra_ip_counts
    combined_mac_counts = global_host_mac + global_vm_mac
    if extra_mac_counts:
        combined_mac_counts += extra_mac_counts

    duplicate_ip_values = {value for value, count in combined_ip_counts.items() if count > 1}
    duplicate_mac_values = {value for value, count in combined_mac_counts.items() if count > 1}
    for summary in summaries:
        _reconcile_duplicate_network_counts(summary, duplicate_ip_values, duplicate_mac_values)
    return duplicate_ip_values, duplicate_mac_values


def build_combined_summary(
    summaries: Iterable[WorkbookSummary],
    source_dir: Path,
    duplicate_ip_values: set[str],
    duplicate_mac_values: set[str],
) -> WorkbookSummary:
    aggregate = WorkbookSummary(path=source_dir, source_vcenter="Combined")
    for summary in summaries:
        aggregate.merge(summary)
    aggregate.host_esxi_versions = int(len(aggregate.host_esxi_version_counts))
    aggregate.host_vendors = int(len(aggregate.host_vendor_counts))
    aggregate.host_models = int(len(aggregate.host_model_counts))
    aggregate.duplicate_uuid_vms = sum(1 for count in aggregate.vm_uuid_counts.values() if count > 1)
    # duplicate_ip_values/duplicate_mac_values come from
    # apply_global_duplicate_network_counts() so they already include Avi's
    # contribution (if any), not just what this run's workbooks contain.
    _reconcile_duplicate_network_counts(aggregate, duplicate_ip_values, duplicate_mac_values)
    return aggregate


def apply_global_duplicate_uuid_counts(summaries: list[WorkbookSummary]) -> None:
    global_uuid_counts: Counter[str] = Counter()
    for summary in summaries:
        global_uuid_counts.update(summary.vm_uuid_counts)

    duplicate_uuid_values = {uuid for uuid, count in global_uuid_counts.items() if count > 1}
    for summary in summaries:
        summary.duplicate_uuid_vms = sum(
            1 for uuid in summary.vm_uuid_counts if uuid in duplicate_uuid_values
        )


def build_avi_summary_frame(avi_network: AviNetworkSummary) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "avi_service_engines_total": avi_network.service_engines_total,
                "avi_vsvips_total": avi_network.vsvips_total,
                "avi_virtual_services_total": avi_network.virtual_services_total,
                "avi_network_subnets_total": avi_network.network_subnets_total,
                "avi_se_with_duplicate_ip": avi_network.se_with_duplicate_ip,
                "avi_vip_with_duplicate_ip": avi_network.vip_with_duplicate_ip,
            }
        ]
    )


def export_summaries(
    workbooks: list[Path],
    summaries: list[WorkbookSummary],
    combined_summary: WorkbookSummary | None,
    output_dir: Path,
    avi_network: AviNetworkSummary | None = None,
) -> tuple[Path, Path, Path | None, list[Path]]:
    controller_keys = sorted(
        {
            controller
            for summary in summaries
            for controller in summary.controller_type_counts
        }
        | (
            set(combined_summary.controller_type_counts)
            if combined_summary is not None
            else set()
        )
    )
    adapter_keys = sorted(
        {
            adapter
            for summary in summaries
            for adapter in summary.network_adapter_type_counts
        }
        | (
            set(combined_summary.network_adapter_type_counts)
            if combined_summary is not None
            else set()
        )
    )
    rows = [summary_to_row(summary, controller_keys, adapter_keys) for summary in summaries]
    if combined_summary is not None:
        rows.append(summary_to_row(combined_summary, controller_keys, adapter_keys))

    report_frame = pd.DataFrame(rows)
    csv_path = output_dir / "rvtools-summary-report.csv"
    cluster_csv_path = output_dir / "rvtools-cluster-summary-report.csv"
    excel_path = output_dir / "rvtools-summary-report.xlsx"
    detail_csv_paths: list[Path] = []

    report_frame.to_csv(csv_path, index=False)
    cluster_report_frame = build_cluster_summary_report(workbooks)
    cluster_report_frame.to_csv(cluster_csv_path, index=False)
    detail_frames = collect_detail_sheets(workbooks, avi_network)
    if avi_network is not None:
        if "service_engines" in avi_network.raw_reports:
            detail_frames["avi_service_engines"] = avi_network.raw_reports["service_engines"]
        if "vsvips" in avi_network.raw_reports:
            detail_frames["avi_vsvips"] = avi_network.raw_reports["vsvips"]
        if "virtual_services" in avi_network.raw_reports:
            detail_frames["avi_virtual_services"] = avi_network.raw_reports["virtual_services"]
        if "network_subnets" in avi_network.raw_reports:
            detail_frames["avi_network_subnets"] = avi_network.raw_reports["network_subnets"]
    detail_frames["host_cpu_models"] = build_counter_breakdown_sheet(
        summaries,
        combined_summary,
        "host_cpu_model_counts",
    )
    detail_frames["host_esxi_versions"] = build_counter_breakdown_sheet(
        summaries,
        combined_summary,
        "host_esxi_version_counts",
    )
    detail_frames["host_vendors"] = build_counter_breakdown_sheet(
        summaries,
        combined_summary,
        "host_vendor_counts",
    )
    detail_frames["host_models"] = build_counter_breakdown_sheet(
        summaries,
        combined_summary,
        "host_model_counts",
    )
    avi_summary_frame = build_avi_summary_frame(avi_network) if avi_network is not None else None

    if OPENPYXL_AVAILABLE:
        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            report_frame.to_excel(writer, index=False, sheet_name="Summary")
            cluster_report_frame.to_excel(writer, index=False, sheet_name="ClusterSummary")
            if avi_summary_frame is not None:
                avi_summary_frame.to_excel(writer, index=False, sheet_name="AviSummary")
            for sheet_name, detail_frame in detail_frames.items():
                detail_frame.to_excel(writer, index=False, sheet_name=excel_sheet_name(sheet_name))
            add_summary_header_links(writer.book, report_frame.columns.tolist())
            add_detail_sheet_back_links(writer.book, report_frame.columns.tolist(), detail_frames.keys())
        return csv_path, cluster_csv_path, excel_path, detail_csv_paths

    if avi_summary_frame is not None:
        avi_summary_csv_path = output_dir / "avi-summary.csv"
        avi_summary_frame.to_csv(avi_summary_csv_path, index=False)
        detail_csv_paths.append(avi_summary_csv_path)

    for sheet_name, detail_frame in detail_frames.items():
        detail_csv_path = output_dir / f"{sheet_name}.csv"
        detail_frame.to_csv(detail_csv_path, index=False)
        detail_csv_paths.append(detail_csv_path)

    return csv_path, cluster_csv_path, None, detail_csv_paths


def summary_fields(summary: WorkbookSummary) -> list[tuple[str, object]]:
    fields = base_summary_fields(summary)
    fields.extend(controller_adapter_fields(summary))
    fields.extend(counter_breakdown_fields(summary))
    return fields


def aggregate_fields(
    aggregate: WorkbookSummary,
    file_count: int,
    source_dir: Path,
) -> list[tuple[str, object]]:
    source_vcenter_field, *rest_fields = base_summary_fields(aggregate)
    fields: list[tuple[str, object]] = [
        ("source_directory", source_dir),
        source_vcenter_field,
        ("files_processed", file_count),
        *rest_fields,
    ]
    fields.extend(controller_adapter_fields(aggregate))
    fields.extend(counter_breakdown_fields(aggregate))
    return fields


def print_summary(summary: WorkbookSummary, heading: str) -> None:
    print(heading)
    print_fields(summary_fields(summary))
    print()


def print_aggregate(aggregate: WorkbookSummary, file_count: int, source_dir: Path) -> None:
    print("Combined summary")
    print_fields(aggregate_fields(aggregate, file_count, source_dir))


def log_summary(logger: logging.Logger, summary: WorkbookSummary, heading: str) -> None:
    logger.info(heading, extra={"console": False})
    for line in format_fields(summary_fields(summary)):
        logger.info(line, extra={"console": False})


def log_aggregate(
    logger: logging.Logger,
    aggregate: WorkbookSummary,
    file_count: int,
    source_dir: Path,
) -> None:
    logger.info("Combined summary", extra={"console": False})
    for line in format_fields(aggregate_fields(aggregate, file_count, source_dir)):
        logger.info(line, extra={"console": False})


def main(logger: logging.Logger) -> int:
    args = parse_args()
    input_dir = Path(args.input_dir).expanduser().resolve()
    report_dir = Path(args.report_dir).expanduser().resolve()
    console_log_path = Path.cwd() / CONSOLE_LOG_BASENAME
    logger.info("Run started for version %s", SCRIPT_VERSION)
    logger.info("Input directory: %s", input_dir)
    logger.info("Report directory: %s", report_dir)

    if not input_dir.exists():
        logger.error("Input directory does not exist: %s", input_dir)
        return 1

    if not input_dir.is_dir():
        logger.error("Input path is not a directory: %s", input_dir)
        return 1

    if not report_dir.exists():
        try:
            report_dir.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            logger.error("Failed to create report directory %s: %s", report_dir, exc)
            return 1
        logger.info("Created report directory: %s", report_dir)

    if not report_dir.is_dir():
        logger.error("Report path is not a directory: %s", report_dir)
        return 1

    avi_network: AviNetworkSummary | None = None
    if args.avi_data:
        avi_dir = Path(args.avi_data).expanduser().resolve()
        if not avi_dir.exists() or not avi_dir.is_dir():
            logger.error("Avi data directory does not exist or is not a directory: %s", avi_dir)
            return 1
        avi_reports = load_avi_reports(avi_dir)
        if not avi_reports:
            logger.warning(
                "No recognized Avi report CSVs found in: %s (expected filenames containing %s)",
                avi_dir,
                ", ".join(sorted(set(AVI_REPORT_FILE_PATTERNS.values()))),
            )
        avi_network = build_avi_network_summary(avi_reports)
        logger.info(
            "Loaded Avi data from %s: %d service engine(s), %d VsVip(s), %d virtual service(s), "
            "%d network subnet row(s)",
            avi_dir,
            avi_network.service_engines_total,
            avi_network.vsvips_total,
            avi_network.virtual_services_total,
            avi_network.network_subnets_total,
        )

    workbooks = discover_workbooks(input_dir)
    if not workbooks:
        logger.error("No XLSX files found in: %s", input_dir)
        return 1
    logger.info("Discovered %d workbook(s) to process", len(workbooks))

    summaries: list[WorkbookSummary] = []
    successful_workbooks: list[Path] = []
    failed_workbooks: list[Path] = []
    for index, workbook in enumerate(workbooks, start=1):
        logger.info("Processing workbook %d/%d: %s", index, len(workbooks), workbook.name)
        try:
            summary = summarize_workbook(workbook)
        except Exception as exc:  # pragma: no cover
            logger.exception(
                "Failed to process %s; skipping this workbook and continuing with the rest: %s",
                workbook.name,
                exc,
            )
            failed_workbooks.append(workbook)
            continue
        successful_workbooks.append(workbook)
        if summary is not None:
            summaries.append(summary)
            logger.info(
                "Processed %s successfully: source_vcenter=%s, vms_total=%d, hosts_total=%d, clusters_total=%d",
                workbook.name,
                summary.source_vcenter,
                summary.vm_total,
                summary.host_total,
                summary.cluster_total,
            )
        else:
            logger.warning("Skipped %s because no usable vInfo sheet was found", workbook.name)

    if failed_workbooks:
        logger.warning(
            "Skipped %d of %d workbook(s) due to processing errors: %s",
            len(failed_workbooks),
            len(workbooks),
            ", ".join(workbook.name for workbook in failed_workbooks),
        )

    if not summaries:
        logger.error("No RVTools workbooks with a vInfo sheet found in: %s", input_dir)
        return 1

    apply_global_duplicate_uuid_counts(summaries)
    # Avi entities only ever contribute IPs (SE data IPs + VIPs), never MACs
    # -- see AviNetworkSummary for why.
    avi_ip_counts = (avi_network.se_ip_counts + avi_network.vip_counts) if avi_network is not None else None
    duplicate_ip_values, duplicate_mac_values = apply_global_duplicate_network_counts(
        summaries,
        avi_ip_counts,
        None,
    )
    if avi_network is not None:
        avi_network.se_with_duplicate_ip = sum(
            count for value, count in avi_network.se_ip_counts.items() if value in duplicate_ip_values
        )
        avi_network.vip_with_duplicate_ip = sum(
            count for value, count in avi_network.vip_counts.items() if value in duplicate_ip_values
        )
    logger.info(
        "Applied global duplicate UUID/IP/MAC reconciliation across %d summary record(s)",
        len(summaries),
    )
    # Only now -- after reconciliation -- are the duplicate UUID/IP/MAC
    # fields on each per-workbook summary correct, since those are resolved
    # across all workbooks together rather than while a single workbook is
    # being summarized.
    for summary in summaries:
        log_summary(logger, summary, f"Individual summary for {summary.source_vcenter}")

    combined_summary: WorkbookSummary | None = None
    if len(summaries) > 1:
        combined_summary = build_combined_summary(summaries, input_dir, duplicate_ip_values, duplicate_mac_values)
        logger.info("Built combined summary across %d workbook(s)", len(summaries))
    else:
        logger.info("Single workbook run detected; consolidated summary will use the only workbook summary")

    if args.summary:
        for summary in summaries:
            print_summary(summary, "Individual summary")

    if combined_summary is not None:
        print_aggregate(combined_summary, len(summaries), input_dir)
        log_aggregate(logger, combined_summary, len(summaries), input_dir)
    else:
        print_summary(summaries[0], "Consolidated summary")
        log_summary(logger, summaries[0], "Consolidated summary")

    if avi_network is not None:
        avi_fields: list[tuple[str, object]] = [
            ("avi_service_engines_total", avi_network.service_engines_total),
            ("avi_vsvips_total", avi_network.vsvips_total),
            ("avi_virtual_services_total", avi_network.virtual_services_total),
            ("avi_network_subnets_total", avi_network.network_subnets_total),
            ("avi_se_with_duplicate_ip", avi_network.se_with_duplicate_ip),
            ("avi_vip_with_duplicate_ip", avi_network.vip_with_duplicate_ip),
        ]
        print()
        print("Avi summary")
        print_fields(avi_fields)
        logger.info("Avi summary", extra={"console": False})
        for line in format_fields(avi_fields):
            logger.info(line, extra={"console": False})

    logger.info("Exporting summary reports")
    csv_path, cluster_csv_path, excel_path, detail_csv_paths = export_summaries(
        successful_workbooks,
        summaries,
        combined_summary,
        report_dir,
        avi_network,
    )
    combined_inventory_csv_paths = export_combined_inventory_csvs(successful_workbooks, report_dir)
    if len(successful_workbooks) <= 1:
        logger.info("Skipped combined vHost/vHBA/vNIC CSV exports because only one workbook was found")
    elif combined_inventory_csv_paths:
        logger.info(
            "Created %d combined inventory CSV export(s) for vHost/vHBA/vNIC",
            len(combined_inventory_csv_paths),
        )
    else:
        logger.info("No combined vHost/vHBA/vNIC CSV exports were created because those sheets were not present")
    logger.info("Finished exporting summary reports")
    print()
    final_fields: list[tuple[str, object]] = [
        ("script_version", SCRIPT_VERSION),
        ("summary_csv_written_to", csv_path),
        ("cluster_summary_csv_written_to", cluster_csv_path),
    ]
    if failed_workbooks:
        final_fields.append(("workbooks_skipped_due_to_errors", len(failed_workbooks)))
    if avi_network is not None:
        final_fields.append(("avi_data_dir", Path(args.avi_data).expanduser().resolve()))
    if excel_path is not None:
        final_fields.append(("summary_excel_written_to", excel_path))
    else:
        final_fields.append(("summary_excel_written_to", "skipped_openpyxl_not_available"))
        for detail_csv_path in detail_csv_paths:
            final_fields.append(("detail_csv_written_to", detail_csv_path))
    for combined_inventory_csv_path in combined_inventory_csv_paths:
        final_fields.append(("combined_inventory_csv_written_to", combined_inventory_csv_path))
    final_fields.append(("console_log_written_to", console_log_path))
    print_fields(final_fields)
    for line in format_fields(final_fields):
        logger.info(line, extra={"console": False})
    logger.info("Run completed successfully")

    return 0


if __name__ == "__main__":
    log_path = Path.cwd() / CONSOLE_LOG_BASENAME
    logger = configure_logger(log_path)
    run_started = datetime.now().isoformat(timespec="seconds")
    logger.info("=== run_started %s version=%s ===", run_started, SCRIPT_VERSION)
    exit_code = 1
    try:
        exit_code = main(logger)
    finally:
        run_finished = datetime.now().isoformat(timespec="seconds")
        logger.info(
            "=== run_finished %s version=%s exit_code=%s ===",
            run_finished,
            SCRIPT_VERSION,
            exit_code,
        )

    raise SystemExit(exit_code)
